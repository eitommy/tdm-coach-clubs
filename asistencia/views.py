from datetime import datetime, timedelta
from calendar import monthrange

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, F
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from django.urls import reverse
from clubs.models import PerfilUsuario, TurnoClub
from clubs.utils import (
    obtener_club_usuario,
    requerir_admin_club,
    requerir_club,
)

from .forms import (
    CategoriaJugadorForm,
    CategoriaEjercicioForm,
    EjercicioForm,
    EntrenamientoInfoForm,
    JugadorForm,
    PerfilForm,
    TrabajoTurnoForm,
    ObservacionJugadorForm,
    MotivoAusenciaForm,
    NoEntrenamientoForm,
    SetPartidoFormSet,
    PartidoTurnoForm,
)

from .models import (
    Asistencia,
    CategoriaJugador,
    CategoriaEjercicio,
    Ejercicio,
    EjercicioRealizado,
    EjercicioTurno,
    Entrenamiento,
    Jugador,
    TrabajoTurno,
    ObservacionJugador,
    PartidoTurno,
    PagoJugador,
)
def _viene_de_onboarding(request):
    return (
        request.GET.get("onboarding") == "1"
        or request.POST.get("onboarding") == "1"
    )
def obtener_o_crear_entrenamiento(
    fecha,
    turno,
    club,
    turno_config=None,
):
    """
    Obtiene o crea el entrenamiento comercial usando TurnoClub.

    `turno` se conserva temporalmente como argumento de compatibilidad
    para resolver el TurnoClub por su orden, pero los registros nuevos
    ya no escriben el campo numérico legacy.
    """
    if turno_config is None:
        turno_config = get_object_or_404(
            TurnoClub,
            club=club,
            dia_semana=fecha.weekday(),
            orden=int(turno),
            activo=True,
        )

    entrenamiento, _ = Entrenamiento.objects.get_or_create(
        club=club,
        fecha=fecha,
        turno_config=turno_config,
    )

    return entrenamiento

def nombre_entrenador(entrenamiento):
    if entrenamiento.responsable_usuario:
        return (
            entrenamiento.responsable_usuario.get_full_name()
            or entrenamiento.responsable_usuario.username
        )

    return "Sin entrenador"


def nombre_turno(entrenamiento, incluir_horario=False):
    if incluir_horario:
        return entrenamiento.nombre_turno_completo

    return entrenamiento.nombre_turno


def asignar_entrenador_si_vacio(entrenamiento, user):
    # Campo antiguo. Ya no se asigna automáticamente.
    # El entrenador real se elige desde la lista de entrenadores del turno.
    return


def turno_bloqueado(entrenamiento):
    return entrenamiento.finalizado or entrenamiento.no_se_entreno


def usuario_puede_operar_club(user, club):
    """
    Verificación explícita para acciones sensibles del entrenamiento.

    Exige:
    - User activo en Django.
    - PerfilUsuario activo.
    - pertenencia al mismo club.
    - rol administrador o entrenador.
    """
    if not user.is_authenticated or not user.is_active:
        return False

    return PerfilUsuario.objects.filter(
        usuario=user,
        club=club,
        activo=True,
        rol__in=[
            PerfilUsuario.Rol.ADMIN,
            PerfilUsuario.Rol.ENTRENADOR,
        ],
    ).exists()


def bloquear_si_usuario_no_operativo(request, club):
    if usuario_puede_operar_club(
        request.user,
        club,
    ):
        return None

    messages.error(
        request,
        (
            "Tu usuario no tiene permisos activos para "
            "operar entrenamientos de este club."
        ),
    )

    return redirect("inicio")


def redirigir_turno_bloqueado(request, entrenamiento):
    messages.warning(
        request,
        "Este turno está bloqueado. Reabrilo o quitá la marca de no entrenamiento para modificarlo.",
    )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )


def obtener_turno_id_para_url(entrenamiento):
    """
    En la parte comercial todo Entrenamiento con club tiene turno_config.
    Los registros legacy sin club no tienen una URL de TurnoClub.
    """
    return entrenamiento.turno_config_id


def redirect_dia_turno(entrenamiento, volver_a=""):
    turno_id = obtener_turno_id_para_url(entrenamiento)

    if turno_id is None:
        return redirect("inicio")

    url = (
        f"/dia/{entrenamiento.fecha.isoformat()}/"
        f"turno/{turno_id}/"
    )

    if volver_a:
        url += volver_a

    return redirect(url)


@login_required
def inicio(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    viene_de_onboarding = request.GET.get("onboarding") == "1"

    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        except ValueError:
            fecha = timezone.localdate()
    else:
        fecha = timezone.localdate()

    tiene_turnos_club = TurnoClub.objects.filter(
        club=club,
        activo=True,
    ).exists()

    tiene_jugadores = Jugador.objects.filter(
        club=club,
        activo=True,
    ).exists()

    tiene_entrenadores = PerfilUsuario.objects.filter(
        club=club,
        rol=PerfilUsuario.Rol.ENTRENADOR,
        activo=True,
        usuario__is_active=True,
    ).exists()

    tiene_ejercicios = Ejercicio.objects.filter(
        club=club,
        activo=True,
    ).exists()

    configuracion_incompleta = not all([
        tiene_turnos_club,
        tiene_jugadores,
        tiene_entrenadores,
        tiene_ejercicios,
    ])

    turnos_configurados = (
        TurnoClub.objects
        .filter(
            club=club,
            dia_semana=fecha.weekday(),
            activo=True,
        )
        .order_by("orden", "hora_inicio", "id")
    )

    turnos_info = []
    total_cargados_dia = 0
    total_marcados_dia = 0
    total_no_entrenados_dia = 0
    total_finalizados_dia = 0

    for turno_config in turnos_configurados:
        entrenamiento = obtener_o_crear_entrenamiento(
            fecha,
            turno_config.orden,
            club,
            turno_config=turno_config,
        )

        cantidad_jugadores = Asistencia.objects.filter(
            entrenamiento=entrenamiento
        ).count()

        cantidad_marcados = (
            Asistencia.objects
            .filter(entrenamiento=entrenamiento)
            .exclude(estado="pendiente")
            .count()
        )

        if not entrenamiento.no_se_entreno:
            total_cargados_dia += cantidad_jugadores
            total_marcados_dia += cantidad_marcados

        if entrenamiento.no_se_entreno:
            total_no_entrenados_dia += 1

        if entrenamiento.finalizado:
            total_finalizados_dia += 1

        if entrenamiento.no_se_entreno:
            estado_texto = "No se entrenó"
            estado_clase = "danger"
            estado_detalle = (
                entrenamiento.get_motivo_no_entrenamiento_display()
            )

            if entrenamiento.detalle_no_entrenamiento:
                estado_detalle = (
                    f"{estado_detalle} — "
                    f"{entrenamiento.detalle_no_entrenamiento}"
                )

        elif entrenamiento.finalizado:
            estado_texto = "Finalizado"
            estado_clase = "success"
            estado_detalle = "Turno cerrado correctamente."

        elif cantidad_jugadores == 0:
            estado_texto = "Sin cargar"
            estado_clase = "secondary"
            estado_detalle = "Todavía no hay jugadores cargados."

        elif cantidad_marcados < cantidad_jugadores:
            estado_texto = "Datos pendientes"
            estado_clase = "warning"
            estado_detalle = (
                f"{cantidad_marcados} de "
                f"{cantidad_jugadores} asistencias marcadas."
            )

        else:
            estado_texto = "Asistencias marcadas"
            estado_clase = "primary"
            estado_detalle = "Falta revisar cierre del turno."

        turnos_info.append({
            "turno_id": turno_config.id,
            "turno_config": turno_config,
            "nombre": turno_config.nombre,
            "hora_inicio": turno_config.hora_inicio,
            "hora_fin": turno_config.hora_fin,
            "cantidad_jugadores": cantidad_jugadores,
            "cantidad_marcados": cantidad_marcados,
            "entrenador": nombre_entrenador(entrenamiento),
            "observaciones": entrenamiento.observaciones,
            "no_se_entreno": entrenamiento.no_se_entreno,
            "motivo_no_entrenamiento": (
                entrenamiento.get_motivo_no_entrenamiento_display()
                if entrenamiento.no_se_entreno
                else ""
            ),
            "detalle_no_entrenamiento": (
                entrenamiento.detalle_no_entrenamiento
            ),
            "finalizado": entrenamiento.finalizado,
            "estado_texto": estado_texto,
            "estado_clase": estado_clase,
            "estado_detalle": estado_detalle,
        })

    contexto = {
        "club": club,
        "fecha": fecha,
        "hoy": timezone.localdate(),
        "ayer": fecha - timedelta(days=1),
        "maniana": fecha + timedelta(days=1),
        "turnos_info": turnos_info,
        "total_jugadores": Jugador.objects.filter(
            club=club,
            activo=True,
        ).count(),
        "total_ejercicios": Ejercicio.objects.filter(
            club=club,
            activo=True,
        ).count(),
        "total_cargados_dia": total_cargados_dia,
        "total_marcados_dia": total_marcados_dia,
        "total_no_entrenados_dia": total_no_entrenados_dia,
        "total_finalizados_dia": total_finalizados_dia,

        # Estados de configuración para mostrar ayudas claras
        # en la pantalla de carga diaria.
        "tiene_turnos_club": tiene_turnos_club,
        "tiene_jugadores": tiene_jugadores,
        "tiene_entrenadores": tiene_entrenadores,
        "tiene_ejercicios": tiene_ejercicios,
        "configuracion_incompleta": configuracion_incompleta,
        "viene_de_onboarding": viene_de_onboarding,
    }

    return render(
        request,
        "asistencia/inicio.html",
        contexto,
    )


@login_required
def ir_a_fecha_asistencia(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    fecha_str = request.GET.get("fecha")

    if not fecha_str:
        fecha = timezone.localdate()
        fecha_str = fecha.isoformat()
    else:
        try:
            fecha = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha = timezone.localdate()
            fecha_str = fecha.isoformat()

    turno_id = request.GET.get("turno_id")

    turno_config = None

    if turno_id:
        try:
            turno_id = int(turno_id)
        except (TypeError, ValueError):
            turno_id = None

    if turno_id:
        turno_config = TurnoClub.objects.filter(
            id=turno_id,
            club=club,
            dia_semana=fecha.weekday(),
            activo=True,
        ).first()

    if turno_config is None:
        turno_config = (
            TurnoClub.objects
            .filter(
                club=club,
                dia_semana=fecha.weekday(),
                activo=True,
            )
            .order_by(
                "orden",
                "hora_inicio",
                "id",
            )
            .first()
        )

    if turno_config is None:
        messages.info(
            request,
            "No hay turnos configurados para esa fecha.",
        )
        return redirect(
            f"/?fecha={fecha_str}"
        )

    return redirect(
        "dia_turno",
        fecha_str=fecha_str,
        turno_id=turno_config.id,
    )

@login_required
def dia_turno(request, fecha_str, turno_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    try:
        fecha = datetime.strptime(
            fecha_str,
            "%Y-%m-%d",
        ).date()
    except ValueError:
        messages.error(
            request,
            "Fecha inválida.",
        )
        return redirect("inicio")

    turno_config = get_object_or_404(
        TurnoClub,
        id=turno_id,
        club=club,
        dia_semana=fecha.weekday(),
        activo=True,
    )

    turno = turno_config.orden

    entrenamiento = obtener_o_crear_entrenamiento(
        fecha,
        turno,
        club,
        turno_config=turno_config,
    )

    if request.GET.get("onboarding") == "1":
        request.session[
            "onboarding_entrenamiento_id"
        ] = entrenamiento.id

    viene_de_onboarding = (
        request.session.get(
            "onboarding_entrenamiento_id"
        )
        == entrenamiento.id
    )

    asistencias = list(
        Asistencia.objects
        .filter(
            entrenamiento=entrenamiento,
            jugador__club=club,
        )
        .select_related("jugador")
    )

    jugadores_disponibles = (
        Jugador.objects
        .filter(
            club=club,
            activo=True,
        )
        .exclude(
            id__in=[
                asistencia.jugador_id
                for asistencia in asistencias
            ]
        )
        .order_by(
            "apellido",
            "nombre",
        )
    )

    ejercicios_turno = (
        EjercicioTurno.objects
        .filter(
            entrenamiento=entrenamiento,
            ejercicio__club=club,
        )
        .select_related(
            "ejercicio",
            "ejercicio__categoria_config",
        )
        .order_by(
            "ejercicio__categoria_config__orden",
            "ejercicio__categoria_config__nombre",
            "ejercicio__nombre",
        )
    )

    ejercicios_turno_por_categoria = {}

    for ejercicio_turno in ejercicios_turno:
        categoria = ejercicio_turno.ejercicio.nombre_categoria

        ejercicios_turno_por_categoria.setdefault(
            categoria,
            [],
        ).append(
            ejercicio_turno.ejercicio.nombre
        )

    total_ejercicios_turno = ejercicios_turno.count()

    for asistencia in asistencias:
        asistencia.observaciones_turno = (
            ObservacionJugador.objects
            .filter(
                jugador=asistencia.jugador,
                entrenamiento=entrenamiento,
            )
            .select_related("creada_por")
            .order_by("-creada_el")
        )

    trabajos_turno = (
        TrabajoTurno.objects
        .filter(
            entrenamiento=entrenamiento,
            jugador_1__club=club,
        )
        .filter(
            Q(jugador_2__isnull=True)
            | Q(jugador_2__club=club)
        )
        .select_related(
            "jugador_1",
            "jugador_2",
        )
        .order_by(
            "cambio",
            "id",
        )
    )

    jugadores_ocupados_por_cambio = {}

    for trabajo in trabajos_turno:
        cambio_clave = str(trabajo.cambio)

        if cambio_clave not in jugadores_ocupados_por_cambio:
            jugadores_ocupados_por_cambio[cambio_clave] = []

        if trabajo.jugador_1_id:
            jugadores_ocupados_por_cambio[cambio_clave].append(
                trabajo.jugador_1_id
            )

        if trabajo.jugador_2_id:
            jugadores_ocupados_por_cambio[cambio_clave].append(
                trabajo.jugador_2_id
            )

    cambios_resumen = []

    jugadores_que_entrenan_ids = {
        asistencia.jugador_id
        for asistencia in asistencias
        if asistencia.estado in [
            "asistio",
            "tarde",
        ]
    }

    numeros_cambio = sorted(
        set(
            trabajos_turno.values_list(
                "cambio",
                flat=True,
            )
        )
    )

    for numero_cambio in numeros_cambio:
        trabajos_del_cambio = trabajos_turno.filter(
            cambio=numero_cambio
        )

        jugadores_asignados_ids = set()

        for trabajo in trabajos_del_cambio:
            if trabajo.jugador_1_id:
                jugadores_asignados_ids.add(
                    trabajo.jugador_1_id
                )

            if trabajo.jugador_2_id:
                jugadores_asignados_ids.add(
                    trabajo.jugador_2_id
                )

        jugadores_pendientes = [
            asistencia.jugador
            for asistencia in asistencias
            if (
                asistencia.estado in [
                    "asistio",
                    "tarde",
                ]
                and asistencia.jugador_id not in jugadores_asignados_ids
            )
        ]

        cantidad_asignados = len(
            jugadores_asignados_ids
            & jugadores_que_entrenan_ids
        )

        cantidad_total = len(jugadores_que_entrenan_ids)

        cambios_resumen.append({
            "numero": numero_cambio,
            "cantidad_asignados": cantidad_asignados,
            "cantidad_total": cantidad_total,
            "jugadores_pendientes": jugadores_pendientes,
            "completo": (
                cantidad_total > 0
                and cantidad_asignados == cantidad_total
            ),
        })

    trabajos_otros_turnos = (
        TrabajoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha=fecha,
            entrenamiento__turno_config__orden__lt=turno_config.orden,
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "jugador_1",
            "jugador_2",
        )
        .order_by(
            "entrenamiento__turno_config__orden",
            "cambio",
            "id",
        )
    )

    partidos_turno = (
        PartidoTurno.objects
        .filter(
            entrenamiento=entrenamiento,
            jugador_1__club=club,
            jugador_2__club=club,
        )
        .select_related(
            "jugador_1",
            "jugador_2",
        )
        .prefetch_related("sets")
        .order_by("id")
    )

    total_jugadores_turno = len(asistencias)

    total_presentes_turno = sum(
        1
        for asistencia in asistencias
        if asistencia.estado == "asistio"
    )

    total_tardes_turno = sum(
        1
        for asistencia in asistencias
        if asistencia.estado == "tarde"
    )

    total_ausentes_turno = sum(
        1
        for asistencia in asistencias
        if asistencia.estado == "ausente"
    )

    total_pendientes_turno = sum(
        1
        for asistencia in asistencias
        if asistencia.estado == "pendiente"
    )

    ausentes_sin_motivo_turno = sum(
        1
        for asistencia in asistencias
        if (
            asistencia.estado == "ausente"
            and not asistencia.motivo_ausencia
        )
    )

    total_trabajos_turno = trabajos_turno.count()
    total_cambios_turno = len(cambios_resumen)

    cambios_completos_turno = sum(
        1
        for cambio in cambios_resumen
        if cambio["completo"]
    )

    cambios_incompletos_turno = (
        total_cambios_turno
        - cambios_completos_turno
    )

    total_partidos_turno = partidos_turno.count()

    partidos_sin_sets_turno = sum(
        1
        for partido in partidos_turno
        if not partido.sets.all()
    )

    total_observaciones_turno = (
        ObservacionJugador.objects
        .filter(entrenamiento=entrenamiento)
        .count()
    )

    alertas_resumen_turno = []

    if total_jugadores_turno == 0:
        alertas_resumen_turno.append(
            "No hay jugadores cargados."
        )

    if total_pendientes_turno > 0:
        alertas_resumen_turno.append(
            (
                f"Hay {total_pendientes_turno} "
                "jugador(es) con asistencia pendiente."
            )
        )

    if ausentes_sin_motivo_turno > 0:
        alertas_resumen_turno.append(
            (
                f"Hay {ausentes_sin_motivo_turno} "
                "ausencia(s) sin motivo."
            )
        )

    if cambios_incompletos_turno > 0:
        alertas_resumen_turno.append(
            (
                f"Hay {cambios_incompletos_turno} "
                "cambio(s) incompleto(s)."
            )
        )

    if partidos_sin_sets_turno > 0:
        alertas_resumen_turno.append(
            (
                f"Hay {partidos_sin_sets_turno} "
                "partido(s) sin sets."
            )
        )

    if entrenamiento.responsable_usuario is None:
        alertas_resumen_turno.append(
            "El turno no tiene entrenador responsable."
        )

    resumen_turno = {
        "total_jugadores": total_jugadores_turno,
        "presentes": total_presentes_turno,
        "tardes": total_tardes_turno,
        "ausentes": total_ausentes_turno,
        "pendientes": total_pendientes_turno,
        "ausentes_sin_motivo": ausentes_sin_motivo_turno,
        "total_trabajos": total_trabajos_turno,
        "total_cambios": total_cambios_turno,
        "cambios_completos": cambios_completos_turno,
        "cambios_incompletos": cambios_incompletos_turno,
        "total_partidos": total_partidos_turno,
        "partidos_sin_sets": partidos_sin_sets_turno,
        "total_observaciones": total_observaciones_turno,
        "total_ejercicios": total_ejercicios_turno,
        "alertas": alertas_resumen_turno,
        "listo_para_finalizar": (
            total_jugadores_turno > 0
            and total_pendientes_turno == 0
            and ausentes_sin_motivo_turno == 0
            and cambios_incompletos_turno == 0
            and partidos_sin_sets_turno == 0
            and entrenamiento.responsable_usuario is not None
        ),
    }

    turnos_del_dia = (
        TurnoClub.objects
        .filter(
            club=club,
            dia_semana=fecha.weekday(),
            activo=True,
        )
        .order_by(
            "orden",
            "hora_inicio",
            "id",
        )
    )

    turno_config = entrenamiento.turno_config

    otros_turnos_disponibles = turnos_del_dia.exclude(
        id=turno_config.id if turno_config else None
    )

    contexto = {
        "entrenamiento": entrenamiento,
        "entrenamiento_form": EntrenamientoInfoForm(
            instance=entrenamiento,
            club=club,
        ),
        "no_entrenamiento_form": NoEntrenamientoForm(
            instance=entrenamiento
        ),
        "trabajo_form": TrabajoTurnoForm(
            entrenamiento=entrenamiento
        ),
        "ejercicios_turno": ejercicios_turno,
        "ejercicios_turno_por_categoria": ejercicios_turno_por_categoria,
        "trabajos_turno": trabajos_turno,
        "jugadores_ocupados_por_cambio": jugadores_ocupados_por_cambio,
        "cambios_resumen": cambios_resumen,
        "trabajos_otros_turnos": trabajos_otros_turnos,
        "nombre_entrenador": nombre_entrenador(
            entrenamiento
        ),
        "asistencias": asistencias,
        "jugadores_disponibles": jugadores_disponibles,
        "partidos_turno": partidos_turno,
        "resumen_turno": resumen_turno,
        "fecha": fecha,
        "turno_id": turno_config.id,
        "turno_config": turno_config,
        "turno_nombre": (
            turno_config.nombre
            if turno_config
            else f"Turno {turno}"
        ),
        "turnos_del_dia": turnos_del_dia,
        "otros_turnos_disponibles": otros_turnos_disponibles,
        "ayer": fecha - timedelta(days=1),
        "maniana": fecha + timedelta(days=1),
        "hoy": timezone.localdate(),
        "viene_de_onboarding": viene_de_onboarding,
    }

    return render(
        request,
        "asistencia/dia_turno.html",
        contexto,
    )

    

@login_required
@require_POST
def copiar_jugadores_turno(request, fecha_str, turno_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    try:
        fecha = datetime.strptime(
            fecha_str,
            "%Y-%m-%d",
        ).date()
    except ValueError:
        messages.error(
            request,
            "Fecha inválida.",
        )
        return redirect("inicio")

    turno_destino_config = get_object_or_404(
        TurnoClub,
        id=turno_id,
        club=club,
        dia_semana=fecha.weekday(),
        activo=True,
    )

    turno_origen_id = request.POST.get("turno_origen")

    try:
        turno_origen_id = int(turno_origen_id)
    except (TypeError, ValueError):
        messages.error(
            request,
            "Tenés que elegir un turno de origen válido.",
        )
        return redirect(
            "dia_turno",
            fecha_str=fecha_str,
            turno_id=turno_destino_config.id,
        )

    if turno_origen_id == turno_destino_config.id:
        messages.error(
            request,
            "No podés copiar jugadores desde el mismo turno.",
        )
        return redirect(
            "dia_turno",
            fecha_str=fecha_str,
            turno_id=turno_destino_config.id,
        )

    turno_origen_config = get_object_or_404(
        TurnoClub,
        id=turno_origen_id,
        club=club,
        dia_semana=fecha.weekday(),
        activo=True,
    )

    entrenamiento_destino = obtener_o_crear_entrenamiento(
        fecha,
        turno_destino_config.orden,
        club,
        turno_config=turno_destino_config,
    )

    if turno_bloqueado(entrenamiento_destino):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento_destino,
        )

    entrenamiento_origen = Entrenamiento.objects.filter(
        club=club,
        fecha=fecha,
        turno_config=turno_origen_config,
    ).first()

    if not entrenamiento_origen:
        messages.error(
            request,
            (
                f"No existe un entrenamiento cargado para "
                f"{turno_origen_config.nombre} en esa fecha."
            ),
        )
        return redirect_dia_turno(
            entrenamiento_destino,
        )

    asistencias_origen = Asistencia.objects.filter(
        entrenamiento=entrenamiento_origen,
        jugador__club=club,
    ).select_related(
        "jugador",
    )

    total_origen = asistencias_origen.count()

    if total_origen == 0:
        messages.warning(
            request,
            (
                f"{turno_origen_config.nombre} "
                "no tiene jugadores cargados."
            ),
        )
        return redirect_dia_turno(
            entrenamiento_destino,
        )

    copiados = 0
    repetidos = 0

    for asistencia_origen in asistencias_origen:
        _, creado = Asistencia.objects.get_or_create(
            entrenamiento=entrenamiento_destino,
            jugador=asistencia_origen.jugador,
            defaults={
                "estado": "pendiente",
            },
        )

        if creado:
            copiados += 1
        else:
            repetidos += 1

    if copiados > 0:
        messages.success(
            request,
            (
                f"Se copiaron {copiados} jugador/es "
                f"desde {turno_origen_config.nombre}."
            ),
        )
    else:
        messages.info(
            request,
            (
                "No se copiaron jugadores nuevos. "
                f"Ya estaban cargados los {repetidos} jugador/es."
            ),
        )

    return redirect_dia_turno(
        entrenamiento_destino,
    )

@login_required
def editar_trabajo_turno(request, trabajo_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    trabajo = get_object_or_404(
        TrabajoTurno.objects.select_related("entrenamiento"),
        id=trabajo_id,
        entrenamiento__club=club,
    )

    entrenamiento = trabajo.entrenamiento

    if request.method == "POST" and turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    if request.method == "POST":
        form = TrabajoTurnoForm(
            request.POST,
            instance=trabajo,
            entrenamiento=entrenamiento,
        )

        if form.is_valid():
            form.save()

            messages.success(
                request,
                "Trabajo actualizado correctamente.",
            )

            return redirect(
                "dia_turno",
                fecha_str=entrenamiento.fecha.isoformat(),
                turno_id=obtener_turno_id_para_url(entrenamiento),
            )
    else:
        form = TrabajoTurnoForm(
            instance=trabajo,
            entrenamiento=entrenamiento,
        )

    return render(
        request,
        "asistencia/editar_trabajo_turno.html",
        {
            "form": form,
            "trabajo": trabajo,
            "entrenamiento": entrenamiento,
        },
    )



@login_required
@require_POST
def agregar_trabajo_turno(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    form = TrabajoTurnoForm(
        request.POST,
        entrenamiento=entrenamiento,
    )

    if form.is_valid():
        trabajo = form.save(commit=False)
        trabajo.entrenamiento = entrenamiento
        trabajo.save()

        messages.success(
            request,
            f"Trabajo agregado correctamente al cambio {trabajo.cambio}.",
        )
    else:
        errores = []

        for lista_errores in form.errors.values():
            for error in lista_errores:
                errores.append(str(error))

        if errores:
            messages.error(
                request,
                "No se pudo agregar el trabajo: " + " ".join(errores),
            )
        else:
            messages.error(
                request,
                "No se pudo agregar el trabajo. Revisá los datos ingresados.",
            )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )
    
@login_required
def crear_partido_turno(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    if request.method == "POST":
        partido_form = PartidoTurnoForm(
            request.POST,
            entrenamiento=entrenamiento,
        )

        partido_temporal = PartidoTurno(
            entrenamiento=entrenamiento
        )

        sets_formset = SetPartidoFormSet(
            request.POST,
            instance=partido_temporal,
            prefix="sets",
        )

        if partido_form.is_valid() and sets_formset.is_valid():
            with transaction.atomic():
                partido = partido_form.save(commit=False)
                partido.entrenamiento = entrenamiento
                partido.save()

                sets_formset.instance = partido

                sets_guardados = sets_formset.save(
                    commit=False
                )

                numero_set = 1

                for set_partido in sets_guardados:
                    set_partido.partido = partido
                    set_partido.numero = numero_set
                    set_partido.save()

                    numero_set += 1

                for set_eliminado in sets_formset.deleted_objects:
                    if set_eliminado.pk:
                        set_eliminado.delete()

            messages.success(
                request,
                (
                    f"Partido guardado: "
                    f"{partido.jugador_1} "
                    f"{partido.resultado_general} "
                    f"{partido.jugador_2}."
                ),
            )

            return redirect(
                "dia_turno",
                fecha_str=entrenamiento.fecha.isoformat(),
                turno_id=obtener_turno_id_para_url(entrenamiento),
            )
    else:
        partido_form = PartidoTurnoForm(
            entrenamiento=entrenamiento,
        )

        partido_temporal = PartidoTurno(
            entrenamiento=entrenamiento
        )

        sets_formset = SetPartidoFormSet(
            instance=partido_temporal,
            prefix="sets",
        )

    return render(
        request,
        "asistencia/crear_partido_turno.html",
        {
            "entrenamiento": entrenamiento,
            "partido_form": partido_form,
            "sets_formset": sets_formset,
        },
    )


@login_required
@require_POST
def eliminar_trabajo_turno(request, trabajo_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    trabajo = get_object_or_404(
        TrabajoTurno.objects.select_related("entrenamiento"),
        id=trabajo_id,
        entrenamiento__club=club,
    )

    entrenamiento = trabajo.entrenamiento
    numero_cambio = trabajo.cambio

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    trabajo.delete()

    messages.success(
        request,
        f"Trabajo del cambio {numero_cambio} eliminado correctamente.",
    )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )


@login_required
@require_POST
def tomar_turno(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    bloqueo_usuario = bloquear_si_usuario_no_operativo(
        request,
        club,
    )

    if bloqueo_usuario:
        return bloqueo_usuario

    if entrenamiento.responsable_usuario is None:
        entrenamiento.responsable_usuario = request.user
        entrenamiento.save(
            update_fields=["responsable_usuario"],
        )
        messages.success(
            request,
            "Tomaste este turno correctamente.",
        )

    elif entrenamiento.responsable_usuario == request.user:
        messages.info(
            request,
            "Este turno ya está asignado a vos.",
        )

    else:
        messages.warning(
            request,
            "Este turno ya fue tomado por otro entrenador.",
        )

    return redirect(
        "dia_turno",
        fecha_str=entrenamiento.fecha.isoformat(),
        turno_id=obtener_turno_id_para_url(entrenamiento),
    )


@login_required
@require_POST
def guardar_info_entrenamiento(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    form = EntrenamientoInfoForm(
        request.POST,
        instance=entrenamiento,
        club=club,
    )

    if form.is_valid():
        form.save()

        messages.success(
            request,
            "Información del turno guardada correctamente.",
        )
    else:
        messages.error(
            request,
            "No se pudo guardar la información del turno. Revisá los datos.",
        )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )
    
@login_required
@require_POST
def guardar_no_entrenamiento(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if entrenamiento.finalizado:
        messages.warning(
            request,
            "Este turno está finalizado. Reabrilo para modificarlo.",
        )

        return redirect_dia_turno(
            entrenamiento,
            request.POST.get("volver_a", ""),
        )

    accion = request.POST.get("accion")

    if accion == "quitar":
        entrenamiento.no_se_entreno = False
        entrenamiento.motivo_no_entrenamiento = ""
        entrenamiento.detalle_no_entrenamiento = ""

        entrenamiento.save(
            update_fields=[
                "no_se_entreno",
                "motivo_no_entrenamiento",
                "detalle_no_entrenamiento",
            ]
        )

        messages.success(
            request,
            "Se quitó la marca de no entrenamiento. El turno quedó abierto para cargar datos.",
        )

        return redirect_dia_turno(
            entrenamiento,
            request.POST.get("volver_a", ""),
        )

    form = NoEntrenamientoForm(
        request.POST,
        instance=entrenamiento,
    )

    if form.is_valid():
        Asistencia.objects.filter(
            entrenamiento=entrenamiento,
        ).delete()

        TrabajoTurno.objects.filter(
            entrenamiento=entrenamiento,
        ).delete()

        PartidoTurno.objects.filter(
            entrenamiento=entrenamiento,
        ).delete()

        EjercicioTurno.objects.filter(
            entrenamiento=entrenamiento,
        ).delete()

        ObservacionJugador.objects.filter(
            entrenamiento=entrenamiento,
        ).delete()

        turno_sin_entrenamiento = form.save(commit=False)
        turno_sin_entrenamiento.no_se_entreno = True
        turno_sin_entrenamiento.finalizado = False
        turno_sin_entrenamiento.finalizado_el = None
        turno_sin_entrenamiento.finalizado_por = None
        turno_sin_entrenamiento.save()

        messages.success(
            request,
            "El turno fue marcado como no entrenado y quedó bloqueado.",
        )
    else:
        messages.error(
            request,
            "No se pudo guardar. Seleccioná un motivo.",
        )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )


@login_required
@require_POST
def finalizar_turno(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    bloqueo_usuario = bloquear_si_usuario_no_operativo(
        request,
        club,
    )

    if bloqueo_usuario:
        return bloqueo_usuario

    if entrenamiento.finalizado:
        messages.info(
            request,
            "Este turno ya está finalizado.",
        )

        return redirect(
            "dia_turno",
            fecha_str=entrenamiento.fecha.isoformat(),
            turno_id=obtener_turno_id_para_url(entrenamiento),
        )

    if entrenamiento.no_se_entreno:
        messages.warning(
            request,
            "Este turno está marcado como no entrenado. Quitá esa marca si querés cargar datos o finalizarlo como entrenamiento.",
        )

        return redirect(
            "dia_turno",
            fecha_str=entrenamiento.fecha.isoformat(),
            turno_id=obtener_turno_id_para_url(entrenamiento),
        )

    errores = []

    asistencias = (
        Asistencia.objects
        .filter(entrenamiento=entrenamiento)
        .select_related("jugador")
    )

    if not asistencias.exists():
        errores.append(
            "El turno no tiene jugadores cargados."
        )

    jugadores_pendientes = [
        str(asistencia.jugador)
        for asistencia in asistencias
        if asistencia.estado == "pendiente"
    ]

    if jugadores_pendientes:
        errores.append(
            "Falta marcar la asistencia de: "
            + ", ".join(jugadores_pendientes)
            + "."
        )

    ausentes_sin_motivo = [
        str(asistencia.jugador)
        for asistencia in asistencias
        if (
            asistencia.estado == "ausente"
            and not asistencia.motivo_ausencia
        )
    ]

    if ausentes_sin_motivo:
        errores.append(
            "Falta cargar el motivo de ausencia de: "
            + ", ".join(ausentes_sin_motivo)
            + "."
        )

    if entrenamiento.responsable_usuario is None:
        errores.append(
            "El turno no tiene un entrenador responsable."
        )

    jugadores_que_entrenaron_ids = set(
        asistencias
        .filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        )
        .values_list(
            "jugador_id",
            flat=True,
        )
    )

    trabajos = (
        TrabajoTurno.objects
        .filter(entrenamiento=entrenamiento)
        .select_related(
            "jugador_1",
            "jugador_2",
        )
    )

    numeros_cambio = sorted(
        set(
            trabajos.values_list(
                "cambio",
                flat=True,
            )
        )
    )

    for numero_cambio in numeros_cambio:
        trabajos_del_cambio = trabajos.filter(
            cambio=numero_cambio
        )

        jugadores_asignados_ids = set()

        for trabajo in trabajos_del_cambio:
            jugadores_asignados_ids.add(
                trabajo.jugador_1_id
            )

            if trabajo.jugador_2_id:
                jugadores_asignados_ids.add(
                    trabajo.jugador_2_id
                )

        jugadores_faltantes_ids = (
            jugadores_que_entrenaron_ids
            - jugadores_asignados_ids
        )

        if jugadores_faltantes_ids:
            jugadores_faltantes = (
                Jugador.objects
                .filter(id__in=jugadores_faltantes_ids)
                .order_by(
                    "apellido",
                    "nombre",
                )
            )

            errores.append(
                f"El cambio {numero_cambio} está incompleto. "
                f"Faltan: "
                + ", ".join(
                    str(jugador)
                    for jugador in jugadores_faltantes
                )
                + "."
            )

    partidos_sin_sets = (
        PartidoTurno.objects
        .filter(entrenamiento=entrenamiento)
        .annotate(cantidad_sets=Count("sets"))
        .filter(cantidad_sets=0)
        .select_related(
            "jugador_1",
            "jugador_2",
        )
    )

    for partido in partidos_sin_sets:
        errores.append(
            f"El partido {partido.jugador_1} vs "
            f"{partido.jugador_2} no tiene sets cargados."
        )

    if errores:
        for error in errores:
            messages.error(
                request,
                error,
            )

        messages.warning(
            request,
            "El turno no pudo finalizarse. Revisá los datos indicados.",
        )

        return redirect(
            "dia_turno",
            fecha_str=entrenamiento.fecha.isoformat(),
            turno_id=obtener_turno_id_para_url(entrenamiento),
        )

    entrenamiento.finalizado = True
    entrenamiento.finalizado_el = timezone.now()
    entrenamiento.finalizado_por = request.user

    entrenamiento.save(
        update_fields=[
            "finalizado",
            "finalizado_el",
            "finalizado_por",
        ]
    )

    onboarding_entrenamiento_id = (
        request.session.get(
            "onboarding_entrenamiento_id"
        )
    )

    if onboarding_entrenamiento_id == entrenamiento.id:
        request.session.pop(
            "onboarding_entrenamiento_id",
            None,
        )

        messages.success(
            request,
            (
                "¡Primer entrenamiento registrado! "
                "Completaste la configuración inicial del club."
            ),
        )

        return redirect(
            "onboarding_club"
        )

    messages.success(
        request,
        "Turno finalizado correctamente.",
    )

    return redirect(
        "dia_turno",
        fecha_str=entrenamiento.fecha.isoformat(),
        turno_id=obtener_turno_id_para_url(entrenamiento),
    )



@login_required
@require_POST
def reabrir_turno(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    bloqueo_usuario = bloquear_si_usuario_no_operativo(
        request,
        club,
    )

    if bloqueo_usuario:
        return bloqueo_usuario

    if not entrenamiento.finalizado:
        messages.info(
            request,
            "Este turno ya está abierto.",
        )

        return redirect(
            "dia_turno",
            fecha_str=entrenamiento.fecha.isoformat(),
            turno_id=obtener_turno_id_para_url(entrenamiento),
        )

    entrenamiento.finalizado = False
    entrenamiento.finalizado_el = None
    entrenamiento.finalizado_por = None

    entrenamiento.save(
        update_fields=[
            "finalizado",
            "finalizado_el",
            "finalizado_por",
        ]
    )

    messages.success(
        request,
        "El turno fue reabierto correctamente.",
    )

    return redirect(
        "dia_turno",
        fecha_str=entrenamiento.fecha.isoformat(),
        turno_id=obtener_turno_id_para_url(entrenamiento),
    )

@login_required
@require_POST
def guardar_observacion_jugador(request, asistencia_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    asistencia = get_object_or_404(
        Asistencia.objects.select_related(
            "jugador",
            "entrenamiento",
        ),
        id=asistencia_id,
        entrenamiento__club=club,
        jugador__club=club,
    )

    entrenamiento = asistencia.entrenamiento
    jugador = asistencia.jugador

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    form = ObservacionJugadorForm(request.POST)

    if form.is_valid():
        observacion = form.save(commit=False)
        observacion.jugador = jugador
        observacion.entrenamiento = entrenamiento
        observacion.creada_por = request.user
        observacion.save()

        messages.success(
            request,
            f"Observación guardada para {jugador}.",
        )
    else:
        errores = []

        for lista_errores in form.errors.values():
            for error in lista_errores:
                errores.append(str(error))

        messages.error(
            request,
            (
                "No se pudo guardar la observación: "
                + " ".join(errores)
            ),
        )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )




@login_required
@require_POST
def agregar_jugador(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    jugadores_ids = request.POST.getlist("jugadores_ids")

    if not jugadores_ids:
        messages.info(
            request,
            "Seleccioná al menos un jugador para agregar al turno.",
        )

        return redirect(
            "dia_turno",
            fecha_str=entrenamiento.fecha.strftime("%Y-%m-%d"),
            turno_id=obtener_turno_id_para_url(entrenamiento),
        )

    jugadores = (
        Jugador.objects
        .filter(
            id__in=jugadores_ids,
            club=club,
            activo=True,
        )
        .order_by(
            "apellido",
            "nombre",
        )
    )

    agregados = 0
    ya_cargados = 0

    for jugador in jugadores:
        asistencia, creado = Asistencia.objects.get_or_create(
            entrenamiento=entrenamiento,
            jugador=jugador,
            defaults={
                "estado": "pendiente",
            },
        )

        if creado:
            agregados += 1
        else:
            ya_cargados += 1

    if agregados > 0:
        messages.success(
            request,
            f"Se agregaron {agregados} jugador(es) a {nombre_turno(entrenamiento)}.",
        )

    if ya_cargados > 0:
        messages.info(
            request,
            f"{ya_cargados} jugador(es) ya estaban cargados en este turno.",
        )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )


@login_required
@require_POST
def copiar_lista_ayer(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    entrenamiento_anterior = (
        Entrenamiento.objects
        .filter(
            club=club,
            fecha__lt=entrenamiento.fecha,
            turno_id=obtener_turno_id_para_url(entrenamiento),
            turno_config=entrenamiento.turno_config,
            asistencias__isnull=False
        )
        .annotate(total_jugadores=Count("asistencias"))
        .filter(total_jugadores__gt=0)
        .order_by("-fecha")
        .first()
    )

    if not entrenamiento_anterior:
        messages.info(request, "No se encontró una lista anterior para copiar.")
        return redirect(
        "dia_turno",
        fecha_str=entrenamiento.fecha.strftime("%Y-%m-%d"),
        turno_id=obtener_turno_id_para_url(entrenamiento),
    )

    asistencias_anteriores = Asistencia.objects.filter(
        entrenamiento=entrenamiento_anterior,
        jugador__club=club,
    ).select_related("jugador")

    jugadores_copiados = 0

    for asistencia_anterior in asistencias_anteriores:
        _, creado = Asistencia.objects.get_or_create(
            entrenamiento=entrenamiento,
            jugador=asistencia_anterior.jugador,
            defaults={"estado": "pendiente"}
        )

        if creado:
            jugadores_copiados += 1

    if jugadores_copiados > 0:
        messages.success(
            request,
            f"Lista anterior copiada correctamente. Se agregaron {jugadores_copiados} jugador/es."
        )
    else:
        messages.info(
            request,
            "La lista anterior ya estaba cargada en este turno."
        )

    return redirect(
        "dia_turno",
        fecha_str=entrenamiento.fecha.strftime("%Y-%m-%d"),
        turno_id=obtener_turno_id_para_url(entrenamiento),
    )


@login_required
@require_POST
def marcar_todos_asistieron(request, entrenamiento_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )
    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    asignar_entrenador_si_vacio(entrenamiento, request.user)

    Asistencia.objects.filter(entrenamiento=entrenamiento).update(estado="asistio")

    messages.success(request, "Todos los jugadores quedaron como asistieron.")
    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )


@login_required
@require_POST
def quitar_jugador(request, asistencia_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    asistencia = get_object_or_404(
        Asistencia.objects.select_related(
            "entrenamiento",
            "jugador",
        ),
        id=asistencia_id,
        entrenamiento__club=club,
        jugador__club=club,
    )

    entrenamiento = asistencia.entrenamiento
    jugador = asistencia.jugador

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    asignar_entrenador_si_vacio(
        entrenamiento,
        request.user,
    )

    trabajos_eliminados, _ = (
        TrabajoTurno.objects
        .filter(entrenamiento=entrenamiento)
        .filter(
            Q(jugador_1=jugador)
            | Q(jugador_2=jugador)
        )
        .delete()
    )

    asistencia.delete()

    if trabajos_eliminados > 0:
        messages.success(
            request,
            (
                f"{jugador} fue quitado del turno y también se eliminaron "
                f"sus trabajos o parejas cargadas."
            ),
        )
    else:
        messages.success(
            request,
            f"{jugador} fue quitado del turno.",
        )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )


@login_required
@require_POST
def cambiar_estado(request, asistencia_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        return JsonResponse(
            {
                "ok": False,
                "error": "Tu usuario no está asociado a un club activo.",
            },
            status=403,
        )

    asistencia = get_object_or_404(
        Asistencia.objects.select_related(
            "entrenamiento",
            "jugador",
        ),
        id=asistencia_id,
        entrenamiento__club=club,
        jugador__club=club,
    )

    if turno_bloqueado(asistencia.entrenamiento):
        return JsonResponse({
            "ok": False,
            "error": "Este turno está bloqueado.",
        }, status=403)

    asignar_entrenador_si_vacio(asistencia.entrenamiento, request.user)

    estado_nuevo = request.POST.get("estado")
    estados_validos = {
        "asistio",
        "ausente",
        "tarde",
    }

    if estado_nuevo not in estados_validos:
        return JsonResponse(
            {
                "ok": False,
                "error": "Estado de asistencia inválido.",
            },
            status=400,
        )

    if asistencia.estado == estado_nuevo:
        asistencia.estado = "pendiente"
    else:
        asistencia.estado = estado_nuevo

    if asistencia.estado != "ausente":
        asistencia.motivo_ausencia = ""
        asistencia.detalle_ausencia = ""

    asistencia.save(
        update_fields=[
            "estado",
            "motivo_ausencia",
            "detalle_ausencia",
        ]
    )

    return JsonResponse({
        "ok": True,
        "estado": asistencia.estado,
    })
    
@login_required
@require_POST
def guardar_motivo_ausencia(request, asistencia_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    asistencia = get_object_or_404(
        Asistencia.objects.select_related(
            "jugador",
            "entrenamiento",
        ),
        id=asistencia_id,
        entrenamiento__club=club,
        jugador__club=club,
    )

    if turno_bloqueado(asistencia.entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            asistencia.entrenamiento,
        )

    form = MotivoAusenciaForm(
        request.POST,
        instance=asistencia,
    )

    if asistencia.estado != "ausente":
        messages.error(
            request,
            "Solo podés cargar un motivo cuando el jugador está ausente.",
        )
    elif form.is_valid():
        form.save()

        messages.success(
            request,
            f"Motivo de ausencia guardado para {asistencia.jugador}.",
        )
    else:
        errores = []

        for lista_errores in form.errors.values():
            for error in lista_errores:
                errores.append(str(error))

        messages.error(
            request,
            "No se pudo guardar el motivo: " + " ".join(errores),
        )

    return redirect_dia_turno(
        asistencia.entrenamiento,
        request.POST.get("volver_a", ""),
    )


@login_required
def lista_jugadores(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    hoy = timezone.localdate()
    inicio_mes = hoy.replace(day=1)

    nombres_meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre",
    }

    try:
        mes_cuota = int(
            request.GET.get(
                "mes",
                hoy.month,
            )
        )
    except (TypeError, ValueError):
        mes_cuota = hoy.month

    if mes_cuota not in nombres_meses:
        mes_cuota = hoy.month

    try:
        anio_cuota = int(
            request.GET.get(
                "anio",
                hoy.year,
            )
        )
    except (TypeError, ValueError):
        anio_cuota = hoy.year

    if anio_cuota < 2000 or anio_cuota > 2100:
        anio_cuota = hoy.year

    inicio_periodo_cuota = datetime(
        anio_cuota,
        mes_cuota,
        1,
    ).date()

    fin_periodo_cuota = datetime(
        anio_cuota,
        mes_cuota,
        monthrange(
            anio_cuota,
            mes_cuota,
        )[1],
    ).date()

    anios_disponibles = list(
        range(
            hoy.year + 1,
            hoy.year - 5,
            -1,
        )
    )

    if anio_cuota not in anios_disponibles:
        anios_disponibles.append(anio_cuota)
        anios_disponibles.sort(reverse=True)

    meses_disponibles = [
        {
            "numero": numero,
            "nombre": nombre,
        }
        for numero, nombre in nombres_meses.items()
    ]

    jugadores = (
        Jugador.objects
        .filter(
            club=club,
            fecha_alta__lte=fin_periodo_cuota,
        )
        .filter(
            Q(fecha_baja__isnull=True)
            | Q(fecha_baja__gte=inicio_periodo_cuota)
        )
        .select_related("categoria")
        .order_by(
            "apellido",
            "nombre",
        )
    )

    pagos_mes = {
        pago.jugador_id: pago
        for pago in (
            PagoJugador.objects
            .filter(
                jugador__club=club,
                anio=anio_cuota,
                mes=mes_cuota,
            )
            .select_related("jugador")
        )
    }

    total_cuotas_periodo = jugadores.count()

    cuotas_pagadas_periodo = sum(
        1
        for jugador_id in jugadores.values_list(
            "id",
            flat=True,
        )
        if (
            pagos_mes.get(jugador_id)
            and pagos_mes[jugador_id].pagado
        )
    )

    cuotas_pendientes_periodo = (
        total_cuotas_periodo
        - cuotas_pagadas_periodo
    )

    porcentaje_cobranza_periodo = (
        round(
            (
                cuotas_pagadas_periodo
                / total_cuotas_periodo
            ) * 100,
            1,
        )
        if total_cuotas_periodo
        else 0
    )

    jugadores_info = []

    for jugador in jugadores:
        ultima_asistencia = (
            Asistencia.objects
            .filter(
                jugador=jugador,
                entrenamiento__club=club,
                entrenamiento__no_se_entreno=False,
            )
            .select_related(
                "entrenamiento",
                "entrenamiento__turno_config",
            )
            .order_by(
                "-entrenamiento__fecha",
                "-entrenamiento__turno_config__orden",
            )
            .first()
        )

        asistencias_mes = (
            Asistencia.objects
            .filter(
                jugador=jugador,
                entrenamiento__club=club,
                entrenamiento__fecha__range=[
                    inicio_mes,
                    hoy,
                ],
                entrenamiento__no_se_entreno=False,
            )
        )

        mes_total = asistencias_mes.count()

        mes_presentes = asistencias_mes.filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        ).count()

        mes_ausentes = asistencias_mes.filter(
            estado="ausente"
        ).count()

        mes_porcentaje = (
            round(
                (
                    mes_presentes
                    / mes_total
                ) * 100,
                1,
            )
            if mes_total
            else 0
        )

        pago_mes = pagos_mes.get(
            jugador.id
        )

        jugadores_info.append({
            "jugador": jugador,
            "ultima_asistencia": ultima_asistencia,
            "mes_total": mes_total,
            "mes_presentes": mes_presentes,
            "mes_ausentes": mes_ausentes,
            "mes_porcentaje": mes_porcentaje,
            "pago_mes": pago_mes,
            "cuota_pagada": bool(
                pago_mes
                and pago_mes.pagado
            ),
        })

    contexto = {
        "jugadores_info": jugadores_info,
        "total_jugadores": jugadores.count(),
        "jugadores_activos": jugadores.filter(
            activo=True,
        ).count(),
        "jugadores_inactivos": jugadores.filter(
            activo=False,
        ).count(),
        "hoy": hoy,
        "inicio_mes": inicio_mes,
        "mes_cuota_nombre": nombres_meses[mes_cuota],
        "mes_cuota_numero": mes_cuota,
        "anio_cuota": anio_cuota,
        "meses_disponibles": meses_disponibles,
        "anios_disponibles": anios_disponibles,
        "inicio_periodo_cuota": inicio_periodo_cuota,
        "fin_periodo_cuota": fin_periodo_cuota,
        "total_cuotas_periodo": total_cuotas_periodo,
        "cuotas_pagadas_periodo": cuotas_pagadas_periodo,
        "cuotas_pendientes_periodo": cuotas_pendientes_periodo,
        "porcentaje_cobranza_periodo": porcentaje_cobranza_periodo,
    }

    return render(
        request,
        "asistencia/lista_jugadores.html",
        contexto,
    )


@requerir_admin_club
def resumen_cuotas(request):
    club = obtener_club_usuario(request.user)
    hoy = timezone.localdate()

    nombres_meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre",
    }

    try:
        mes = int(
            request.GET.get(
                "mes",
                hoy.month,
            )
        )
    except (TypeError, ValueError):
        mes = hoy.month

    if mes not in nombres_meses:
        mes = hoy.month

    try:
        anio = int(
            request.GET.get(
                "anio",
                hoy.year,
            )
        )
    except (TypeError, ValueError):
        anio = hoy.year

    if anio < 2000 or anio > 2100:
        anio = hoy.year

    inicio_periodo = datetime(
        anio,
        mes,
        1,
    ).date()

    fin_periodo = datetime(
        anio,
        mes,
        monthrange(
            anio,
            mes,
        )[1],
    ).date()

    jugadores_periodo = (
        Jugador.objects
        .filter(
            club=club,
            fecha_alta__lte=fin_periodo,
        )
        .filter(
            Q(fecha_baja__isnull=True)
            | Q(fecha_baja__gte=inicio_periodo)
        )
        .select_related("categoria")
        .order_by(
            "apellido",
            "nombre",
        )
    )

    pagos_periodo = {
        pago.jugador_id: pago
        for pago in (
            PagoJugador.objects
            .filter(
                jugador__club=club,
                anio=anio,
                mes=mes,
            )
        )
    }

    total_periodo = jugadores_periodo.count()

    pagados_periodo = sum(
        1
        for jugador in jugadores_periodo
        if (
            pagos_periodo.get(jugador.id)
            and pagos_periodo[jugador.id].pagado
        )
    )

    pendientes_periodo = (
        total_periodo
        - pagados_periodo
    )

    porcentaje_periodo = (
        round(
            (
                pagados_periodo
                / total_periodo
            ) * 100,
            1,
        )
        if total_periodo
        else 0
    )

    jugadores_pendientes_periodo = []

    for jugador in jugadores_periodo:
        pago = pagos_periodo.get(
            jugador.id
        )

        if not (
            pago
            and pago.pagado
        ):
            jugadores_pendientes_periodo.append({
                "jugador": jugador,
                "categoria": (
                    jugador.categoria.nombre
                    if jugador.categoria
                    else "Sin categoría"
                ),
            })

    def inicio_mes(fecha):
        return fecha.replace(day=1)

    def sumar_meses(fecha, cantidad):
        indice = (
            fecha.year * 12
            + fecha.month
            - 1
            + cantidad
        )

        nuevo_anio = indice // 12
        nuevo_mes = indice % 12 + 1

        return datetime(
            nuevo_anio,
            nuevo_mes,
            1,
        ).date()

    corte_seleccionado = inicio_periodo

    corte_hoy = hoy.replace(
        day=1,
    )

    corte_deuda = min(
        corte_seleccionado,
        corte_hoy,
    )

    jugadores_para_deuda = (
        Jugador.objects
        .filter(
            club=club,
            fecha_alta__lte=(
                corte_deuda.replace(
                    day=monthrange(
                        corte_deuda.year,
                        corte_deuda.month,
                    )[1]
                )
            ),
        )
        .select_related("categoria")
        .order_by(
            "apellido",
            "nombre",
        )
    )

    pagos_hasta_corte = (
        PagoJugador.objects
        .filter(
            jugador__club=club,
        )
        .filter(
            Q(anio__lt=corte_deuda.year)
            | Q(
                anio=corte_deuda.year,
                mes__lte=corte_deuda.month,
            )
        )
    )

    pagos_por_jugador = {}

    for pago in pagos_hasta_corte:
        pagos_por_jugador.setdefault(
            pago.jugador_id,
            {},
        )[
            (pago.anio, pago.mes)
        ] = pago

    deuda_jugadores = []

    for jugador in jugadores_para_deuda:
        inicio = inicio_mes(
            jugador.fecha_alta
        )

        fin = corte_deuda

        if jugador.fecha_baja:
            fin = min(
                fin,
                inicio_mes(
                    jugador.fecha_baja
                ),
            )

        if inicio > fin:
            continue

        fecha_cursor = inicio

        meses_correspondientes = 0
        meses_pagados = 0
        meses_pendientes = []

        while fecha_cursor <= fin:
            meses_correspondientes += 1

            pago = pagos_por_jugador.get(
                jugador.id,
                {},
            ).get(
                (
                    fecha_cursor.year,
                    fecha_cursor.month,
                )
            )

            if pago and pago.pagado:
                meses_pagados += 1
            else:
                meses_pendientes.append({
                    "anio": fecha_cursor.year,
                    "mes": fecha_cursor.month,
                    "mes_nombre": nombres_meses[
                        fecha_cursor.month
                    ],
                })

            fecha_cursor = sumar_meses(
                fecha_cursor,
                1,
            )

        if meses_pendientes:
            deuda_jugadores.append({
                "jugador": jugador,
                "categoria": (
                    jugador.categoria.nombre
                    if jugador.categoria
                    else "Sin categoría"
                ),
                "meses_correspondientes": (
                    meses_correspondientes
                ),
                "meses_pagados": (
                    meses_pagados
                ),
                "meses_pendientes": len(
                    meses_pendientes
                ),
                "detalle_pendientes": list(
                    reversed(
                        meses_pendientes
                    )
                ),
            })

    deuda_jugadores.sort(
        key=lambda item: (
            -item["meses_pendientes"],
            item["jugador"].apellido.lower(),
            item["jugador"].nombre.lower(),
        )
    )

    jugadores_con_deuda = len(
        deuda_jugadores
    )

    cuotas_pendientes_acumuladas = sum(
        item["meses_pendientes"]
        for item in deuda_jugadores
    )

    jugadores_reincidentes = sum(
        1
        for item in deuda_jugadores
        if item["meses_pendientes"] >= 2
    )

    evolucion = []

    for desplazamiento in range(-5, 1):
        fecha_mes = sumar_meses(
            inicio_periodo,
            desplazamiento,
        )

        inicio_mes_evolucion = fecha_mes

        fin_mes_evolucion = datetime(
            fecha_mes.year,
            fecha_mes.month,
            monthrange(
                fecha_mes.year,
                fecha_mes.month,
            )[1],
        ).date()

        jugadores_mes = (
            Jugador.objects
            .filter(
                club=club,
                fecha_alta__lte=fin_mes_evolucion,
            )
            .filter(
                Q(fecha_baja__isnull=True)
                | Q(
                    fecha_baja__gte=(
                        inicio_mes_evolucion
                    )
                )
            )
        )

        total_mes = jugadores_mes.count()

        ids_jugadores_mes = list(
            jugadores_mes.values_list(
                "id",
                flat=True,
            )
        )

        pagados_mes = (
            PagoJugador.objects
            .filter(
                jugador_id__in=ids_jugadores_mes,
                anio=fecha_mes.year,
                mes=fecha_mes.month,
                pagado=True,
            )
            .count()
        )

        porcentaje_mes = (
            round(
                (
                    pagados_mes
                    / total_mes
                ) * 100,
                1,
            )
            if total_mes
            else 0
        )

        evolucion.append({
            "anio": fecha_mes.year,
            "mes": fecha_mes.month,
            "mes_nombre": nombres_meses[
                fecha_mes.month
            ],
            "total": total_mes,
            "pagados": pagados_mes,
            "pendientes": (
                total_mes
                - pagados_mes
            ),
            "porcentaje": porcentaje_mes,
        })

    meses_disponibles = [
        {
            "numero": numero,
            "nombre": nombre,
        }
        for numero, nombre
        in nombres_meses.items()
    ]

    anios_disponibles = list(
        range(
            hoy.year + 1,
            hoy.year - 5,
            -1,
        )
    )

    if anio not in anios_disponibles:
        anios_disponibles.append(
            anio
        )
        anios_disponibles.sort(
            reverse=True
        )

    contexto = {
        "mes": mes,
        "anio": anio,
        "mes_nombre": nombres_meses[mes],
        "meses_disponibles": meses_disponibles,
        "anios_disponibles": anios_disponibles,

        "total_periodo": total_periodo,
        "pagados_periodo": pagados_periodo,
        "pendientes_periodo": pendientes_periodo,
        "porcentaje_periodo": porcentaje_periodo,
        "jugadores_pendientes_periodo": (
            jugadores_pendientes_periodo
        ),

        "jugadores_con_deuda": jugadores_con_deuda,
        "cuotas_pendientes_acumuladas": (
            cuotas_pendientes_acumuladas
        ),
        "jugadores_reincidentes": (
            jugadores_reincidentes
        ),
        "deuda_jugadores": deuda_jugadores,

        "evolucion": evolucion,
        "corte_deuda": corte_deuda,
    }

    return render(
        request,
        "asistencia/resumen_cuotas.html",
        contexto,
    )


@requerir_admin_club
def exportar_cuotas_excel(request):
    club = obtener_club_usuario(request.user)

    hoy = timezone.localdate()

    nombres_meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre",
    }

    try:
        mes_cuota = int(
            request.GET.get(
                "mes",
                hoy.month,
            )
        )
    except (TypeError, ValueError):
        mes_cuota = hoy.month

    if mes_cuota not in nombres_meses:
        mes_cuota = hoy.month

    try:
        anio_cuota = int(
            request.GET.get(
                "anio",
                hoy.year,
            )
        )
    except (TypeError, ValueError):
        anio_cuota = hoy.year

    if anio_cuota < 2000 or anio_cuota > 2100:
        anio_cuota = hoy.year

    inicio_periodo = datetime(
        anio_cuota,
        mes_cuota,
        1,
    ).date()

    fin_periodo = datetime(
        anio_cuota,
        mes_cuota,
        monthrange(
            anio_cuota,
            mes_cuota,
        )[1],
    ).date()

    jugadores = (
        Jugador.objects
        .filter(
            club=club,
            fecha_alta__lte=fin_periodo,
        )
        .filter(
            Q(fecha_baja__isnull=True)
            | Q(fecha_baja__gte=inicio_periodo)
        )
        .select_related("categoria")
        .order_by(
            "apellido",
            "nombre",
        )
    )

    pagos = {
        pago.jugador_id: pago
        for pago in (
            PagoJugador.objects
            .filter(
                jugador__club=club,
                anio=anio_cuota,
                mes=mes_cuota,
            )
            .select_related("jugador")
        )
    }

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Cuotas"

    hoja.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=7,
    )

    celda_titulo = hoja.cell(
        row=1,
        column=1,
        value=(
            f"Cuotas · {club.nombre} · "
            f"{nombres_meses[mes_cuota]} {anio_cuota}"
        ),
    )
    celda_titulo.font = Font(
        bold=True,
        size=14,
    )
    celda_titulo.alignment = Alignment(
        horizontal="center",
    )

    encabezados = [
        "Jugador",
        "Categoría",
        "Estado jugador",
        "Fecha de alta",
        "Fecha de baja",
        "Estado cuota",
        "Fecha de pago",
    ]

    fila_encabezado = 3

    for columna, encabezado in enumerate(
        encabezados,
        start=1,
    ):
        celda = hoja.cell(
            row=fila_encabezado,
            column=columna,
            value=encabezado,
        )
        celda.font = Font(
            bold=True,
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    fila = fila_encabezado + 1

    total = 0
    pagados = 0
    pendientes = 0

    for jugador in jugadores:
        pago = pagos.get(
            jugador.id
        )

        cuota_pagada = bool(
            pago
            and pago.pagado
        )

        if cuota_pagada:
            estado_cuota = "Pagado"
            pagados += 1
        else:
            estado_cuota = "Pendiente"
            pendientes += 1

        total += 1

        valores = [
            str(jugador),
            (
                jugador.categoria.nombre
                if jugador.categoria
                else "Sin categoría"
            ),
            (
                "Activo"
                if jugador.activo
                else "Inactivo"
            ),
            jugador.fecha_alta,
            jugador.fecha_baja,
            estado_cuota,
            (
                pago.fecha_pago
                if pago
                else None
            ),
        ]

        for columna, valor in enumerate(
            valores,
            start=1,
        ):
            celda = hoja.cell(
                row=fila,
                column=columna,
                value=valor,
            )

            if columna in [4, 5, 7] and valor:
                celda.number_format = "DD/MM/YYYY"

        fila += 1

    fila_resumen = fila + 2

    hoja.cell(
        row=fila_resumen,
        column=1,
        value="Resumen",
    ).font = Font(
        bold=True,
        size=12,
    )

    hoja.cell(
        row=fila_resumen + 1,
        column=1,
        value="Jugadores del período",
    )
    hoja.cell(
        row=fila_resumen + 1,
        column=2,
        value=total,
    )

    hoja.cell(
        row=fila_resumen + 2,
        column=1,
        value="Pagados",
    )
    hoja.cell(
        row=fila_resumen + 2,
        column=2,
        value=pagados,
    )

    hoja.cell(
        row=fila_resumen + 3,
        column=1,
        value="Pendientes",
    )
    hoja.cell(
        row=fila_resumen + 3,
        column=2,
        value=pendientes,
    )

    porcentaje = (
        round(
            (
                pagados
                / total
            ) * 100,
            1,
        )
        if total
        else 0
    )

    hoja.cell(
        row=fila_resumen + 4,
        column=1,
        value="Cobranza",
    )
    hoja.cell(
        row=fila_resumen + 4,
        column=2,
        value=porcentaje / 100,
    ).number_format = "0.0%"

    anchos = {
        1: 28,
        2: 20,
        3: 18,
        4: 16,
        5: 16,
        6: 16,
        7: 16,
    }

    for numero_columna, ancho in anchos.items():
        hoja.column_dimensions[
            get_column_letter(
                numero_columna
            )
        ].width = ancho

    hoja.freeze_panes = "A4"
    hoja.auto_filter.ref = (
        f"A3:G{max(fila - 1, 3)}"
    )

    response = HttpResponse(
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    nombre_archivo = (
        "cuotas_"
        f"{anio_cuota}_"
        f"{mes_cuota:02d}.xlsx"
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{nombre_archivo}"'
    )

    workbook.save(response)

    return response


@requerir_admin_club
@require_POST
def cambiar_estado_cuota(request, jugador_id):
    club = obtener_club_usuario(request.user)

    jugador = get_object_or_404(
        Jugador,
        id=jugador_id,
        club=club,
    )

    hoy = timezone.localdate()

    try:
        mes_cuota = int(
            request.POST.get(
                "mes",
                hoy.month,
            )
        )
    except (TypeError, ValueError):
        mes_cuota = hoy.month

    if mes_cuota < 1 or mes_cuota > 12:
        mes_cuota = hoy.month

    try:
        anio_cuota = int(
            request.POST.get(
                "anio",
                hoy.year,
            )
        )
    except (TypeError, ValueError):
        anio_cuota = hoy.year

    if anio_cuota < 2000 or anio_cuota > 2100:
        anio_cuota = hoy.year

    inicio_periodo_cuota = datetime(
        anio_cuota,
        mes_cuota,
        1,
    ).date()

    fin_periodo_cuota = datetime(
        anio_cuota,
        mes_cuota,
        monthrange(
            anio_cuota,
            mes_cuota,
        )[1],
    ).date()

    jugador_pertenecia_al_club = (
        jugador.fecha_alta <= fin_periodo_cuota
        and (
            jugador.fecha_baja is None
            or jugador.fecha_baja >= inicio_periodo_cuota
        )
    )

    if not jugador_pertenecia_al_club:
        messages.error(
            request,
            (
                "No podés modificar la cuota de ese período "
                "porque el jugador no pertenecía al club."
            ),
        )

        return redirect(
            f"{reverse('lista_jugadores')}"
            f"?mes={mes_cuota}&anio={anio_cuota}"
        )

    pago, _ = PagoJugador.objects.get_or_create(
        jugador=jugador,
        anio=anio_cuota,
        mes=mes_cuota,
    )

    accion = request.POST.get("accion")

    if accion == "pagado":
        pago.pagado = True
        pago.fecha_pago = hoy

        messages.success(
            request,
            f"Cuota de {jugador} marcada como pagada.",
        )

    elif accion == "pendiente":
        pago.pagado = False
        pago.fecha_pago = None

        messages.info(
            request,
            f"Cuota de {jugador} marcada como pendiente.",
        )

    else:
        messages.error(
            request,
            "Estado de cuota inválido.",
        )

        return redirect(
            f"{reverse('lista_jugadores')}"
            f"?mes={mes_cuota}&anio={anio_cuota}"
        )

    pago.save(
        update_fields=[
            "pagado",
            "fecha_pago",
        ]
    )

    return redirect(
        f"{reverse('lista_jugadores')}"
        f"?mes={mes_cuota}&anio={anio_cuota}"
    )


@requerir_club
def crear_jugador(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    viene_de_onboarding = _viene_de_onboarding(
        request
    )

    if request.method == "POST":
        form = JugadorForm(
            request.POST,
            club=club,
        )

        if form.is_valid():
            jugador = form.save(commit=False)
            jugador.club = club
            jugador.save()

            messages.success(
                request,
                "Jugador agregado correctamente.",
            )

            if viene_de_onboarding:
                return redirect(
                    "onboarding_club"
                )

            return redirect(
                "lista_jugadores"
            )

    else:
        form = JugadorForm(
            club=club,
        )

    return render(
        request,
        "asistencia/form_jugador.html",
        {
            "form": form,
            "titulo": "Agregar jugador",
            "viene_de_onboarding": viene_de_onboarding,
        },
    )

@requerir_club
def editar_jugador(request, pk):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    jugador = get_object_or_404(
        Jugador,
        pk=pk,
        club=club,
    )

    if request.method == "POST":
        form = JugadorForm(
            request.POST,
            instance=jugador,
            club=club,
        )

        if form.is_valid():
            jugador_editado = form.save(commit=False)
            jugador_editado.club = club
            jugador_editado.save()

            messages.success(
                request,
                "Jugador editado correctamente.",
            )
            return redirect("lista_jugadores")
    else:
        form = JugadorForm(
            instance=jugador,
            club=club,
        )

    return render(
        request,
        "asistencia/form_jugador.html",
        {
            "form": form,
            "titulo": "Editar jugador",
        },
    )


@requerir_admin_club
def lista_categorias_jugador(request):
    club = obtener_club_usuario(request.user)

    categorias = list(
        CategoriaJugador.objects
        .filter(club=club)
        .annotate(
            cantidad_jugadores=Count("jugadores")
        )
        .order_by(
            "orden",
            "nombre",
            "id",
        )
    )

    total_categorias = len(categorias)

    categorias_activas = sum(
        1
        for categoria in categorias
        if categoria.activo
    )

    categorias_inactivas = (
        total_categorias
        - categorias_activas
    )

    jugadores_asignados = sum(
        categoria.cantidad_jugadores
        for categoria in categorias
    )

    for indice, categoria in enumerate(categorias):
        categoria.puede_subir = indice > 0
        categoria.puede_bajar = (
            indice < total_categorias - 1
        )

    return render(
        request,
        "asistencia/lista_categorias_jugador.html",
        {
            "categorias": categorias,
            "total_categorias": total_categorias,
            "categorias_activas": categorias_activas,
            "categorias_inactivas": categorias_inactivas,
            "jugadores_asignados": jugadores_asignados,
        },
    )


@requerir_admin_club
@require_POST
def mover_categoria_jugador(
    request,
    pk,
    direccion,
):
    club = obtener_club_usuario(request.user)

    if direccion not in {
        "subir",
        "bajar",
    }:
        messages.error(
            request,
            "Movimiento de categoría inválido.",
        )
        return redirect(
            "lista_categorias_jugador"
        )

    with transaction.atomic():
        categorias = list(
            CategoriaJugador.objects
            .select_for_update()
            .filter(club=club)
            .order_by(
                "orden",
                "nombre",
                "id",
            )
        )

        categoria_actual = None
        indice_actual = None

        for indice, categoria in enumerate(categorias):
            if categoria.pk == pk:
                categoria_actual = categoria
                indice_actual = indice
                break

        if categoria_actual is None:
            messages.error(
                request,
                "La categoría no existe.",
            )
            return redirect(
                "lista_categorias_jugador"
            )

        for indice, categoria in enumerate(
            categorias,
            start=1,
        ):
            if categoria.orden != indice:
                CategoriaJugador.objects.filter(
                    pk=categoria.pk,
                    club=club,
                ).update(
                    orden=indice
                )
                categoria.orden = indice

        if direccion == "subir":
            indice_destino = indice_actual - 1
        else:
            indice_destino = indice_actual + 1

        if (
            indice_destino < 0
            or indice_destino >= len(categorias)
        ):
            return redirect(
                "lista_categorias_jugador"
            )

        categoria_destino = categorias[indice_destino]

        orden_actual = categoria_actual.orden
        orden_destino = categoria_destino.orden

        CategoriaJugador.objects.filter(
            pk=categoria_actual.pk,
            club=club,
        ).update(
            orden=orden_destino
        )

        CategoriaJugador.objects.filter(
            pk=categoria_destino.pk,
            club=club,
        ).update(
            orden=orden_actual
        )

    return redirect(
        "lista_categorias_jugador"
    )


@requerir_admin_club
def crear_categoria_jugador(request):
    club = obtener_club_usuario(request.user)

    if request.method == "POST":
        form = CategoriaJugadorForm(
            request.POST,
            club=club,
        )

        if form.is_valid():
            categoria = form.save(commit=False)
            categoria.club = club
            categoria.save()

            messages.success(
                request,
                "Categoría de jugadores creada correctamente.",
            )

            return redirect(
                "lista_categorias_jugador"
            )

    else:
        form = CategoriaJugadorForm(
            club=club,
        )

    return render(
        request,
        "asistencia/form_categoria_jugador.html",
        {
            "form": form,
            "titulo": "Agregar categoría de jugadores",
        },
    )


@requerir_admin_club
def editar_categoria_jugador(request, pk):
    club = obtener_club_usuario(request.user)

    categoria = get_object_or_404(
        CategoriaJugador,
        pk=pk,
        club=club,
    )

    if request.method == "POST":
        form = CategoriaJugadorForm(
            request.POST,
            instance=categoria,
            club=club,
        )

        if form.is_valid():
            categoria_editada = form.save(commit=False)
            categoria_editada.club = club
            categoria_editada.save()

            messages.success(
                request,
                "Categoría de jugadores editada correctamente.",
            )

            return redirect(
                "lista_categorias_jugador"
            )

    else:
        form = CategoriaJugadorForm(
            instance=categoria,
            club=club,
        )

    return render(
        request,
        "asistencia/form_categoria_jugador.html",
        {
            "form": form,
            "titulo": "Editar categoría de jugadores",
            "categoria": categoria,
        },
    )


@login_required
def lista_ejercicios(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    ejercicios = (
        Ejercicio.objects
        .filter(club=club)
        .select_related("categoria_config")
        .order_by(
            "categoria_config__orden",
            "categoria_config__nombre",
            "nombre",
        )
    )

    return render(
        request,
        "asistencia/lista_ejercicios.html",
        {
            "ejercicios": ejercicios,
        },
    )


@requerir_club
def lista_categorias_ejercicio(request):
    club = obtener_club_usuario(request.user)

    categorias = (
        CategoriaEjercicio.objects
        .filter(club=club)
        .annotate(
            cantidad_ejercicios=Count("ejercicios")
        )
        .order_by(
            "orden",
            "nombre",
        )
    )

    return render(
        request,
        "asistencia/lista_categorias_ejercicio.html",
        {
            "categorias": categorias,
        },
    )


@requerir_club
def crear_categoria_ejercicio(request):
    club = obtener_club_usuario(request.user)

    viene_de_onboarding = _viene_de_onboarding(
        request
    )

    if request.method == "POST":
        form = CategoriaEjercicioForm(
            request.POST,
            club=club,
        )

        if form.is_valid():
            categoria = form.save(commit=False)
            categoria.club = club
            categoria.save()

            messages.success(
                request,
                "Categoría creada correctamente.",
            )

            if viene_de_onboarding:
                return redirect(
                    reverse("crear_ejercicio")
                    + "?onboarding=1"
                )

            return redirect(
                "lista_categorias_ejercicio"
            )

    else:
        form = CategoriaEjercicioForm(
            club=club,
        )

    return render(
        request,
        "asistencia/form_categoria_ejercicio.html",
        {
            "form": form,
            "titulo": "Agregar categoría",
            "viene_de_onboarding": viene_de_onboarding,
        },
    )


@requerir_club
def editar_categoria_ejercicio(request, pk):
    club = obtener_club_usuario(request.user)

    categoria = get_object_or_404(
        CategoriaEjercicio,
        pk=pk,
        club=club,
    )

    if request.method == "POST":
        form = CategoriaEjercicioForm(
            request.POST,
            instance=categoria,
            club=club,
        )

        if form.is_valid():
            categoria_editada = form.save(commit=False)
            categoria_editada.club = club
            categoria_editada.save()

            messages.success(
                request,
                "Categoría editada correctamente.",
            )
            return redirect(
                "lista_categorias_ejercicio"
            )
    else:
        form = CategoriaEjercicioForm(
            instance=categoria,
            club=club,
        )

    return render(
        request,
        "asistencia/form_categoria_ejercicio.html",
        {
            "form": form,
            "titulo": "Editar categoría",
            "categoria": categoria,
        },
    )


@requerir_club
def crear_ejercicio(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    viene_de_onboarding = _viene_de_onboarding(
        request
    )

    if not CategoriaEjercicio.objects.filter(
        club=club,
        activo=True,
    ).exists():

        messages.info(
            request,
            "Primero creá al menos una categoría de ejercicios.",
        )

        if viene_de_onboarding:
            return redirect(
                reverse("crear_categoria_ejercicio")
                + "?onboarding=1"
            )

        return redirect(
            "lista_categorias_ejercicio"
        )

    if request.method == "POST":
        form = EjercicioForm(
            request.POST,
            club=club,
        )

        if form.is_valid():
            ejercicio = form.save(commit=False)
            ejercicio.club = club
            ejercicio.save()

            messages.success(
                request,
                "Ejercicio agregado correctamente.",
            )

            if viene_de_onboarding:
                return redirect(
                    "onboarding_club"
                )

            return redirect(
                "lista_ejercicios"
            )

    else:
        form = EjercicioForm(
            club=club,
        )

    return render(
        request,
        "asistencia/form_ejercicio.html",
        {
            "form": form,
            "titulo": "Agregar ejercicio",
            "viene_de_onboarding": viene_de_onboarding,
        },
    )

@requerir_club
def editar_ejercicio(request, pk):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    ejercicio = get_object_or_404(
        Ejercicio.objects.select_related(
            "categoria_config"
        ),
        pk=pk,
        club=club,
    )

    if request.method == "POST":
        form = EjercicioForm(
            request.POST,
            instance=ejercicio,
            club=club,
        )

        if form.is_valid():
            ejercicio_editado = form.save(commit=False)
            ejercicio_editado.club = club
            ejercicio_editado.save()

            messages.success(
                request,
                "Ejercicio editado correctamente.",
            )
            return redirect("lista_ejercicios")
    else:
        form = EjercicioForm(
            instance=ejercicio,
            club=club,
        )

    return render(
        request,
        "asistencia/form_ejercicio.html",
        {
            "form": form,
            "titulo": "Editar ejercicio",
        },
    )


@login_required
def cargar_ejercicios(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    fecha_str = request.GET.get("fecha")
    turno_id = request.GET.get("turno_id")

    if fecha_str:
        try:
            fecha = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha = timezone.localdate()
    else:
        fecha = timezone.localdate()

    turnos_del_dia = (
        TurnoClub.objects
        .filter(
            club=club,
            activo=True,
            dia_semana=fecha.weekday(),
        )
        .order_by(
            "orden",
            "hora_inicio",
            "id",
        )
    )

    turno_config = None

    if turno_id:
        try:
            turno_id = int(turno_id)
        except (TypeError, ValueError):
            turno_id = None

    if turno_id:
        turno_config = turnos_del_dia.filter(
            id=turno_id,
        ).first()

    if turno_config is None:
        turno_config = turnos_del_dia.first()

    if turno_config is None:
        messages.info(
            request,
            "No hay turnos configurados para ese día.",
        )
        return redirect(
            f"/?fecha={fecha.isoformat()}"
        )

    entrenamiento = obtener_o_crear_entrenamiento(
        fecha,
        turno_config.orden,
        club,
        turno_config=turno_config,
    )

    ejercicios_guardados = list(
        EjercicioTurno.objects
        .filter(
            entrenamiento=entrenamiento,
            ejercicio__club=club,
        )
        .values_list(
            "ejercicio_id",
            flat=True,
        )
    )

    categorias_activas = (
        CategoriaEjercicio.objects
        .filter(
            club=club,
            activo=True,
        )
        .order_by(
            "orden",
            "nombre",
        )
    )

    ejercicios_por_categoria = {}

    for categoria in categorias_activas:
        ejercicios_por_categoria[categoria.nombre] = (
            Ejercicio.objects
            .filter(
                club=club,
                categoria_config=categoria,
                activo=True,
            )
            .order_by("nombre")
        )


    contexto = {
        "entrenamiento": entrenamiento,
        "fecha": fecha,
        "hoy": timezone.localdate(),
        "turno_id": turno_config.id,
        "turno_config": turno_config,
        "turnos_del_dia": turnos_del_dia,
        "ayer": fecha - timedelta(days=1),
        "maniana": fecha + timedelta(days=1),
        "ejercicios_por_categoria": ejercicios_por_categoria,
        "ejercicios_guardados": ejercicios_guardados,
    }

    return render(
        request,
        "asistencia/cargar_ejercicios.html",
        contexto,
    )


@login_required
@require_POST
def guardar_ejercicios(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    entrenamiento_id = request.POST.get("entrenamiento_id")
    ejercicio_ids = request.POST.getlist("ejercicios")

    if not entrenamiento_id:
        messages.error(
            request,
            "No se pudo identificar el turno. Volvé a entrar desde la pantalla de asistencia.",
        )
        return redirect("inicio")

    entrenamiento = get_object_or_404(
        Entrenamiento,
        id=entrenamiento_id,
        club=club,
    )

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    EjercicioTurno.objects.filter(
        entrenamiento=entrenamiento,
    ).delete()

    ejercicios_validos = (
        Ejercicio.objects
        .filter(
            id__in=ejercicio_ids,
            club=club,
            activo=True,
        )
    )

    for ejercicio in ejercicios_validos:
        EjercicioTurno.objects.create(
            entrenamiento=entrenamiento,
            ejercicio=ejercicio,
        )

    messages.success(
        request,
        "Ejercicios del turno guardados correctamente.",
    )

    return redirect_dia_turno(
        entrenamiento,
    )


@login_required
def seguimiento_semanal(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    jugadores = (
        Jugador.objects
        .filter(
            club=club,
            activo=True,
        )
        .order_by("apellido", "nombre")
    )

    jugador_id = request.GET.get("jugador")
    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha_base = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_base = timezone.localdate()
    else:
        fecha_base = timezone.localdate()

    inicio_semana = fecha_base - timedelta(
        days=fecha_base.weekday()
    )

    fin_semana = inicio_semana + timedelta(days=4)

    dias_semana = [
        inicio_semana + timedelta(days=i)
        for i in range(5)
    ]

    semana_anterior = inicio_semana - timedelta(days=7)
    semana_siguiente = inicio_semana + timedelta(days=7)

    jugador_seleccionado = None
    filas = []
    resumen = None

    if jugador_id:
        jugador_seleccionado = (
            Jugador.objects
            .filter(
                id=jugador_id,
                club=club,
                activo=True,
            )
            .first()
        )

        if jugador_seleccionado:
            total_dias_programados = 0
            total_presentes = 0
            total_tardes = 0
            total_ausencias = 0
            total_dias_no_entrenados = 0
            total_ejercicios = 0
            total_trabajos = 0
            total_partidos = 0
            total_partidos_ganados = 0
            total_partidos_perdidos = 0

            for dia in dias_semana:
                entrenamientos_no_entrenados = (
                    Entrenamiento.objects
                    .filter(
                        club=club,
                        fecha=dia,
                        no_se_entreno=True,
                    )
                    .select_related("turno_config")
                    .order_by(
                        "turno_config__orden",
                    )
                )

                turnos_no_entrenados_config = [
                    {
                        "turno_config": entrenamiento_no.turno_config,
                        "nombre": entrenamiento_no.nombre_turno_completo,
                    }
                    for entrenamiento_no in entrenamientos_no_entrenados
                ]

                nombres_turnos_no_entrenados = [
                    item["nombre"]
                    for item in turnos_no_entrenados_config
                ]

                motivos_no_entrenamiento = []

                for entrenamiento_no in entrenamientos_no_entrenados:
                    motivo_texto = (
                        entrenamiento_no.get_motivo_no_entrenamiento_display()
                    )

                    if entrenamiento_no.detalle_no_entrenamiento:
                        motivo_texto = (
                            f"{motivo_texto}: "
                            f"{entrenamiento_no.detalle_no_entrenamiento}"
                        )

                    motivos_no_entrenamiento.append({
                        "turno_config": entrenamiento_no.turno_config,
                        "motivo": motivo_texto,
                    })

                asistencias_dia = (
                    Asistencia.objects
                    .filter(
                        jugador=jugador_seleccionado,
                        entrenamiento__club=club,
                        entrenamiento__fecha=dia,
                    )
                    .select_related(
                        "entrenamiento",
                        "entrenamiento__turno_config",
                    )
                )

                turnos_entrenamientos = [
                    asistencia.entrenamiento
                    for asistencia in asistencias_dia
                ]

                turnos = [
                    entrenamiento.nombre_turno_completo
                    for entrenamiento in turnos_entrenamientos
                ]

                estados = list(
                    asistencias_dia.values_list(
                        "estado",
                        flat=True,
                    )
                )

                if asistencias_dia.exists():
                    total_dias_programados += 1

                if entrenamientos_no_entrenados.exists():
                    total_dias_no_entrenados += 1

                entrenamientos_del_jugador_dia = [
                    asistencia.entrenamiento
                    for asistencia in asistencias_dia
                ]

                ejercicios_qs = (
                    EjercicioTurno.objects
                    .filter(
                        entrenamiento__in=entrenamientos_del_jugador_dia,
                    )
                    .select_related(
                        "ejercicio",
                        "ejercicio__categoria_config",
                        "entrenamiento",
                        "entrenamiento__turno_config",
                    )
                    .order_by(
                        "entrenamiento__turno_config__orden",
                        "ejercicio__categoria_config__orden",
                        "ejercicio__categoria_config__nombre",
                        "ejercicio__nombre",
                    )
                )

                ejercicios_por_turno = []
                turnos_dict = {}

                for item in ejercicios_qs:
                    turno_config = item.entrenamiento.turno_config
                    clave_turno = turno_config.id

                    categoria = item.ejercicio.nombre_categoria

                    if clave_turno not in turnos_dict:
                        turnos_dict[clave_turno] = {
                            "turno_config": turno_config,
                            "categorias": {},
                        }

                    if categoria not in turnos_dict[clave_turno]["categorias"]:
                        turnos_dict[clave_turno]["categorias"][categoria] = []

                    turnos_dict[clave_turno]["categorias"][categoria].append(
                        item.ejercicio.nombre
                    )

                ejercicios_por_turno = list(turnos_dict.values())

                cantidad_ejercicios = ejercicios_qs.count()
                total_ejercicios += cantidad_ejercicios

                trabajos_qs = (
                    TrabajoTurno.objects
                    .filter(
                        entrenamiento__club=club,
                        entrenamiento__fecha=dia,
                    )
                    .filter(
                        Q(jugador_1=jugador_seleccionado)
                        | Q(jugador_2=jugador_seleccionado)
                    )
                    .select_related(
                        "entrenamiento",
                        "entrenamiento__turno_config",
                        "jugador_1",
                        "jugador_2",
                    )
                    .order_by(
                        "entrenamiento__turno_config__orden",
                        "cambio",
                        "id",
                    )
                )

                trabajos_dia = []

                for trabajo in trabajos_qs:
                    if trabajo.tipo == TrabajoTurno.Tipo.PAREJA:
                        if trabajo.jugador_1_id == jugador_seleccionado.id:
                            companero = trabajo.jugador_2
                        else:
                            companero = trabajo.jugador_1

                        descripcion = f"Con {companero}"
                    else:
                        descripcion = trabajo.get_tipo_display()

                    trabajos_dia.append({
                        "turno_config": trabajo.entrenamiento.turno_config,
                        "cambio": trabajo.cambio,
                        "tipo": trabajo.tipo,
                        "tipo_texto": trabajo.get_tipo_display(),
                        "descripcion": descripcion,
                        "detalle": trabajo.detalle,
                    })

                cantidad_trabajos = len(trabajos_dia)
                total_trabajos += cantidad_trabajos

                partidos_qs = (
                    PartidoTurno.objects
                    .filter(
                        entrenamiento__club=club,
                        entrenamiento__fecha=dia,
                    )
                    .filter(
                        Q(jugador_1=jugador_seleccionado)
                        | Q(jugador_2=jugador_seleccionado)
                    )
                    .select_related(
                        "entrenamiento",
                        "entrenamiento__turno_config",
                        "jugador_1",
                        "jugador_2",
                    )
                    .prefetch_related("sets")
                    .order_by(
                        "entrenamiento__turno_config__orden",
                        "id",
                    )
                )

                partidos_dia = []

                for partido in partidos_qs:
                    if partido.jugador_1_id == jugador_seleccionado.id:
                        rival = partido.jugador_2
                        sets_propios = partido.sets_jugador_1
                        sets_rival = partido.sets_jugador_2

                        sets_resultados = [
                            (
                                f"{set_partido.puntos_jugador_1}-"
                                f"{set_partido.puntos_jugador_2}"
                            )
                            for set_partido in partido.sets.all()
                        ]

                    else:
                        rival = partido.jugador_1
                        sets_propios = partido.sets_jugador_2
                        sets_rival = partido.sets_jugador_1

                        sets_resultados = [
                            (
                                f"{set_partido.puntos_jugador_2}-"
                                f"{set_partido.puntos_jugador_1}"
                            )
                            for set_partido in partido.sets.all()
                        ]

                    if partido.ganador == jugador_seleccionado:
                        resultado_clase = "success"
                        resultado_texto = "Victoria"
                        total_partidos_ganados += 1

                    elif partido.ganador:
                        resultado_clase = "danger"
                        resultado_texto = "Derrota"
                        total_partidos_perdidos += 1

                    else:
                        resultado_clase = "secondary"
                        resultado_texto = "Sin definir"

                    partidos_dia.append({
                        "turno_config": partido.entrenamiento.turno_config,
                        "rival": rival,
                        "resultado": f"{sets_propios}-{sets_rival}",
                        "resultado_clase": resultado_clase,
                        "resultado_texto": resultado_texto,
                        "sets": sets_resultados,
                        "detalle": partido.detalle,
                    })

                cantidad_partidos = len(partidos_dia)
                total_partidos += cantidad_partidos

                if "asistio" in estados:
                    asistencia_texto = "Asistió"
                    asistencia_clase = "success"
                    total_presentes += 1

                elif "tarde" in estados:
                    asistencia_texto = "Tarde"
                    asistencia_clase = "warning text-dark"
                    total_presentes += 1
                    total_tardes += 1

                elif "ausente" in estados:
                    asistencia_texto = "Ausente"
                    asistencia_clase = "danger"
                    total_ausencias += 1

                elif asistencias_dia.exists():
                    asistencia_texto = "Sin marcar"
                    asistencia_clase = "secondary"

                elif entrenamientos_no_entrenados.exists():
                    asistencia_texto = "No se entrenó"
                    asistencia_clase = "danger"

                elif (
                    cantidad_ejercicios > 0
                    or cantidad_trabajos > 0
                    or cantidad_partidos > 0
                ):
                    asistencia_texto = "Sin asistencia"
                    asistencia_clase = "secondary"

                else:
                    asistencia_texto = "Sin actividad"
                    asistencia_clase = "light text-dark"

                turnos_combinados = list(
                    dict.fromkeys(
                        turnos + nombres_turnos_no_entrenados
                    )
                )

                filas.append({
                    "dia": dia,
                    "asistencia_texto": asistencia_texto,
                    "asistencia_clase": asistencia_clase,
                    "turnos": (
                        ", ".join(turnos_combinados)
                        if turnos_combinados
                        else "-"
                    ),
                    "no_se_entreno": entrenamientos_no_entrenados.exists(),
                    "turnos_no_entrenados_config": turnos_no_entrenados_config,
                    "motivos_no_entrenamiento": motivos_no_entrenamiento,
                    "ejercicios_por_turno": ejercicios_por_turno,
                    "cantidad_ejercicios": cantidad_ejercicios,
                    "trabajos_dia": trabajos_dia,
                    "cantidad_trabajos": cantidad_trabajos,
                    "partidos_dia": partidos_dia,
                    "cantidad_partidos": cantidad_partidos,
                })

            porcentaje = (
                round(
                    (
                        total_presentes
                        / total_dias_programados
                    ) * 100,
                    1,
                )
                if total_dias_programados
                else 0
            )

            resumen = {
                "total_dias_programados": total_dias_programados,
                "total_presentes": total_presentes,
                "total_tardes": total_tardes,
                "total_ausencias": total_ausencias,
                "total_dias_no_entrenados": total_dias_no_entrenados,
                "total_ejercicios": total_ejercicios,
                "total_trabajos": total_trabajos,
                "total_partidos": total_partidos,
                "total_partidos_ganados": total_partidos_ganados,
                "total_partidos_perdidos": total_partidos_perdidos,
                "porcentaje": porcentaje,
            }

    contexto = {
        "jugadores": jugadores,
        "jugador_seleccionado": jugador_seleccionado,
        "fecha_base": fecha_base,
        "inicio_semana": inicio_semana,
        "fin_semana": fin_semana,
        "semana_anterior": semana_anterior,
        "semana_siguiente": semana_siguiente,
        "filas": filas,
        "resumen": resumen,
    }

    return render(
        request,
        "asistencia/seguimiento_semanal.html",
        contexto,
    )


    

@login_required
def historial_jugador(request, jugador_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    jugador = get_object_or_404(
        Jugador,
        id=jugador_id,
        club=club,
    )

    hoy = timezone.localdate()
    fecha_desde_str = request.GET.get("desde")
    fecha_hasta_str = request.GET.get("hasta")

    if fecha_desde_str:
        try:
            fecha_desde = datetime.strptime(
                fecha_desde_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_desde = hoy - timedelta(days=30)
    else:
        fecha_desde = hoy - timedelta(days=30)

    if fecha_hasta_str:
        try:
            fecha_hasta = datetime.strptime(
                fecha_hasta_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_hasta = hoy
    else:
        fecha_hasta = hoy

    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde

    inicio_mes = fecha_hasta.replace(day=1)

    asistencias = (
        Asistencia.objects
        .filter(
            jugador=jugador,
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
            entrenamiento__no_se_entreno=False,
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__responsable_usuario",
            "entrenamiento__turno_config",
        )
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
        )
    )

    total_asistencias = asistencias.count()

    total_presentes = asistencias.filter(
        Q(estado="asistio")
        | Q(estado="tarde")
    ).count()

    total_tardes = asistencias.filter(
        estado="tarde",
    ).count()

    total_ausencias = asistencias.filter(
        estado="ausente",
    ).count()

    total_pendientes = asistencias.filter(
        estado="pendiente",
    ).count()

    ausencias = asistencias.filter(
        estado="ausente",
    )

    ausencias_justificadas = (
        ausencias
        .exclude(motivo_ausencia="")
        .exclude(motivo_ausencia="sin_aviso")
        .count()
    )

    ausencias_sin_aviso = ausencias.filter(
        Q(motivo_ausencia="")
        | Q(motivo_ausencia="sin_aviso")
    ).count()

    porcentaje_asistencia = (
        round(
            (
                total_presentes
                / total_asistencias
            ) * 100,
            1,
        )
        if total_asistencias
        else 0
    )

    asistencias_mes = asistencias.filter(
        entrenamiento__fecha__range=[
            inicio_mes,
            fecha_hasta,
        ]
    )

    mes_total = asistencias_mes.count()

    mes_presentes = asistencias_mes.filter(
        Q(estado="asistio")
        | Q(estado="tarde")
    ).count()

    mes_ausencias = asistencias_mes.filter(
        estado="ausente",
    ).count()

    mes_tardes = asistencias_mes.filter(
        estado="tarde",
    ).count()

    mes_porcentaje = (
        round(
            (
                mes_presentes
                / mes_total
            ) * 100,
            1,
        )
        if mes_total
        else 0
    )

    for asistencia in asistencias:
        asistencia.ejercicios_turno = (
            EjercicioTurno.objects
            .filter(
                entrenamiento=asistencia.entrenamiento,
            )
            .select_related(
                "ejercicio",
                "ejercicio__categoria_config",
            )
            .order_by(
                "ejercicio__categoria_config__orden",
                "ejercicio__categoria_config__nombre",
                "ejercicio__nombre",
            )
        )

    entrenamientos_del_jugador = [
        asistencia.entrenamiento
        for asistencia in asistencias
    ]

    ejercicios_turno = (
        EjercicioTurno.objects
        .filter(
            entrenamiento__in=entrenamientos_del_jugador,
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "ejercicio",
            "ejercicio__categoria_config",
        )
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "ejercicio__categoria_config__orden",
            "ejercicio__categoria_config__nombre",
            "ejercicio__nombre",
        )
    )

    ejercicios = []

    for item in ejercicios_turno:
        ejercicios.append({
            "fecha": item.entrenamiento.fecha,
            "turno_config": item.entrenamiento.turno_config,
            "ejercicio": item.ejercicio,
        })

    ejercicios_frecuentes = (
        ejercicios_turno
        .values(
            "ejercicio__nombre",
            "ejercicio__categoria_config__nombre",
        )
        .annotate(total=Count("id"))
        .order_by("-total")[:8]
    )

    partidos = (
        PartidoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
        )
        .filter(
            Q(jugador_1=jugador)
            | Q(jugador_2=jugador)
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "jugador_1",
            "jugador_2",
        )
        .prefetch_related("sets")
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "-id",
        )
    )

    total_partidos = partidos.count()
    total_partidos_ganados = 0
    total_partidos_perdidos = 0
    total_partidos_empatados = 0

    for partido in partidos:
        sets_jugador_1 = partido.sets_jugador_1
        sets_jugador_2 = partido.sets_jugador_2

        if sets_jugador_1 == sets_jugador_2:
            total_partidos_empatados += 1
            continue

        if partido.jugador_1_id == jugador.id:
            if sets_jugador_1 > sets_jugador_2:
                total_partidos_ganados += 1
            else:
                total_partidos_perdidos += 1
        else:
            if sets_jugador_2 > sets_jugador_1:
                total_partidos_ganados += 1
            else:
                total_partidos_perdidos += 1

    partidos_definidos = (
        total_partidos_ganados
        + total_partidos_perdidos
    )

    porcentaje_victorias = (
        round(
            (
                total_partidos_ganados
                / partidos_definidos
            ) * 100,
            1,
        )
        if partidos_definidos
        else 0
    )

    trabajos = (
        TrabajoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
        )
        .filter(
            Q(jugador_1=jugador)
            | Q(jugador_2=jugador)
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "jugador_1",
            "jugador_2",
        )
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "cambio",
        )
    )

    observaciones = (
        ObservacionJugador.objects
        .filter(
            jugador=jugador,
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "creada_por",
        )
        .order_by(
            "-entrenamiento__fecha",
            "-creada_el",
        )
    )

    turnos_no_entrenados = (
        Entrenamiento.objects
        .filter(
            club=club,
            fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
            no_se_entreno=True,
        )
        .select_related("turno_config")
        .order_by(
            "-fecha",
            "turno_config__orden",
        )
    )

    # Historial mensual de cuotas del jugador.
    nombres_meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre",
    }

    fecha_inicio_cuotas = jugador.fecha_alta.replace(day=1)

    fecha_fin_membresia = (
        jugador.fecha_baja
        if jugador.fecha_baja
        else hoy
    )

    fecha_fin_cuotas = min(
        fecha_fin_membresia,
        hoy,
    ).replace(day=1)

    pagos_jugador = {
        (pago.anio, pago.mes): pago
        for pago in PagoJugador.objects.filter(
            jugador=jugador,
        )
    }

    historial_cuotas = []

    if fecha_inicio_cuotas <= fecha_fin_cuotas:
        anio = fecha_inicio_cuotas.year
        mes = fecha_inicio_cuotas.month

        while (
            anio < fecha_fin_cuotas.year
            or (
                anio == fecha_fin_cuotas.year
                and mes <= fecha_fin_cuotas.month
            )
        ):
            pago = pagos_jugador.get(
                (anio, mes)
            )

            historial_cuotas.append({
                "anio": anio,
                "mes": mes,
                "mes_nombre": nombres_meses[mes],
                "pago": pago,
                "pagado": bool(
                    pago
                    and pago.pagado
                ),
                "fecha_pago": (
                    pago.fecha_pago
                    if pago
                    else None
                ),
                "observacion": (
                    pago.observacion
                    if pago
                    else ""
                ),
            })

            if mes == 12:
                mes = 1
                anio += 1
            else:
                mes += 1

    historial_cuotas.reverse()

    total_cuotas = len(historial_cuotas)

    cuotas_pagadas = sum(
        1
        for cuota in historial_cuotas
        if cuota["pagado"]
    )

    cuotas_pendientes = (
        total_cuotas
        - cuotas_pagadas
    )

    porcentaje_cuotas_pagadas = (
        round(
            (
                cuotas_pagadas
                / total_cuotas
            ) * 100,
            1,
        )
        if total_cuotas
        else 0
    )

    resumen = {
        "total_asistencias": total_asistencias,
        "total_presentes": total_presentes,
        "total_tardes": total_tardes,
        "total_ausencias": total_ausencias,
        "total_pendientes": total_pendientes,
        "ausencias_justificadas": ausencias_justificadas,
        "ausencias_sin_aviso": ausencias_sin_aviso,
        "porcentaje_asistencia": porcentaje_asistencia,

        "mes_total": mes_total,
        "mes_presentes": mes_presentes,
        "mes_ausencias": mes_ausencias,
        "mes_tardes": mes_tardes,
        "mes_porcentaje": mes_porcentaje,

        "total_partidos": total_partidos,
        "total_partidos_ganados": total_partidos_ganados,
        "total_partidos_perdidos": total_partidos_perdidos,
        "total_partidos_empatados": total_partidos_empatados,
        "porcentaje_victorias": porcentaje_victorias,

        "total_trabajos": trabajos.count(),
        "total_ejercicios": len(ejercicios),
        "total_observaciones": observaciones.count(),
        "turnos_no_entrenados": turnos_no_entrenados.count(),

        "total_cuotas": total_cuotas,
        "cuotas_pagadas": cuotas_pagadas,
        "cuotas_pendientes": cuotas_pendientes,
        "porcentaje_cuotas_pagadas": porcentaje_cuotas_pagadas,
    }

    contexto = {
        "jugador": jugador,
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "inicio_mes": inicio_mes,
        "asistencias": asistencias,
        "partidos": partidos,
        "trabajos": trabajos,
        "ejercicios": ejercicios,
        "ejercicios_frecuentes": ejercicios_frecuentes,
        "observaciones": observaciones,
        "turnos_no_entrenados": turnos_no_entrenados,
        "historial_cuotas": historial_cuotas,
        "resumen": resumen,

        # Compatibilidad con el template viejo, por si alguna parte todavía los usa.
        "total_asistencias": total_asistencias,
        "total_presentes": total_presentes,
        "total_tardes": total_tardes,
        "total_ausencias": total_ausencias,
        "porcentaje_asistencia": porcentaje_asistencia,
        "total_partidos": total_partidos,
        "total_partidos_ganados": total_partidos_ganados,
        "total_partidos_perdidos": total_partidos_perdidos,
        "total_partidos_empatados": total_partidos_empatados,
    }

    return render(
        request,
        "asistencia/historial_jugador.html",
        contexto,
    )


@login_required
def reportes(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    hoy = timezone.localdate()
    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha_base = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_base = hoy
    else:
        fecha_base = hoy

    inicio_semana = fecha_base - timedelta(
        days=fecha_base.weekday()
    )

    fin_semana = inicio_semana + timedelta(days=4)

    if fin_semana > hoy:
        fin_semana = hoy

    inicio_mes = fecha_base.replace(day=1)

    if inicio_mes.month == 12:
        inicio_mes_siguiente = inicio_mes.replace(
            year=inicio_mes.year + 1,
            month=1,
        )
    else:
        inicio_mes_siguiente = inicio_mes.replace(
            month=inicio_mes.month + 1,
        )

    fin_mes = inicio_mes_siguiente - timedelta(days=1)

    if fin_mes > hoy:
        fin_mes = hoy

    semana_anterior = inicio_semana - timedelta(days=7)
    semana_siguiente = inicio_semana + timedelta(days=7)
    mes_anterior = inicio_mes - timedelta(days=1)
    mes_siguiente = inicio_mes_siguiente

    turnos_no_entrenados_semana = (
        Entrenamiento.objects
        .filter(
            club=club,
            fecha__range=[
                inicio_semana,
                fin_semana,
            ],
            no_se_entreno=True,
        )
        .select_related("turno_config")
        .order_by(
            "-fecha",
            "turno_config__orden",
        )
    )

    turnos_no_entrenados_mes = (
        Entrenamiento.objects
        .filter(
            club=club,
            fecha__range=[
                inicio_mes,
                fin_mes,
            ],
            no_se_entreno=True,
        )
        .select_related("turno_config")
        .order_by(
            "-fecha",
            "turno_config__orden",
        )
    )

    total_turnos_no_entrenados_semana = (
        turnos_no_entrenados_semana.count()
    )

    total_turnos_no_entrenados_mes = (
        turnos_no_entrenados_mes.count()
    )

    dias_no_entrenados_semana = (
        turnos_no_entrenados_semana
        .values("fecha")
        .distinct()
        .count()
    )

    dias_no_entrenados_mes = (
        turnos_no_entrenados_mes
        .values("fecha")
        .distinct()
        .count()
    )

    motivos_no_entrenamiento_mes = (
        turnos_no_entrenados_mes
        .values("motivo_no_entrenamiento")
        .annotate(total=Count("id"))
        .order_by("-total")
    )

    motivos_no_entrenamiento = []

    opciones_motivos = dict(
        Entrenamiento.MotivoNoEntrenamiento.choices
    )

    for item in motivos_no_entrenamiento_mes:
        motivo_codigo = item["motivo_no_entrenamiento"]

        if motivo_codigo:
            motivo_texto = opciones_motivos.get(
                motivo_codigo,
                motivo_codigo,
            )
        else:
            motivo_texto = "Sin motivo"

        motivos_no_entrenamiento.append({
            "motivo": motivo_texto,
            "total": item["total"],
        })

    datos = []

    total_jugadores = 0
    total_semana_presentes = 0
    total_semana_ausentes = 0
    total_mes_presentes = 0
    total_mes_ausentes = 0
    total_mes_pendientes = 0

    for jugador in (
        Jugador.objects
        .filter(
            club=club,
            activo=True,
        )
        .order_by(
            "apellido",
            "nombre",
        )
    ):
        total_jugadores += 1

        asistencias_semana = Asistencia.objects.filter(
            jugador=jugador,
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_semana,
                fin_semana,
            ],
            entrenamiento__no_se_entreno=False,
        )

        asistencias_mes = Asistencia.objects.filter(
            jugador=jugador,
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
            entrenamiento__no_se_entreno=False,
        )

        semana_total = asistencias_semana.count()

        semana_presentes = asistencias_semana.filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        ).count()

        semana_tardes = asistencias_semana.filter(
            estado="tarde"
        ).count()

        semana_ausentes = asistencias_semana.filter(
            estado="ausente"
        ).count()

        semana_pendientes = asistencias_semana.filter(
            estado="pendiente"
        ).count()

        semana_porcentaje = (
            round(
                (
                    semana_presentes
                    / semana_total
                ) * 100,
                1,
            )
            if semana_total
            else 0
        )

        mes_total = asistencias_mes.count()

        mes_presentes = asistencias_mes.filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        ).count()

        mes_tardes = asistencias_mes.filter(
            estado="tarde"
        ).count()

        ausencias_mes = asistencias_mes.filter(
            estado="ausente"
        )

        mes_ausentes = ausencias_mes.count()

        mes_pendientes = asistencias_mes.filter(
            estado="pendiente"
        ).count()

        mes_ausencias_justificadas = (
            ausencias_mes
            .exclude(motivo_ausencia="")
            .exclude(motivo_ausencia="sin_aviso")
            .count()
        )

        mes_ausencias_sin_aviso = ausencias_mes.filter(
            Q(motivo_ausencia="sin_aviso")
            | Q(motivo_ausencia="")
        ).count()

        motivo_mas_frecuente = (
            ausencias_mes
            .exclude(motivo_ausencia="")
            .values("motivo_ausencia")
            .annotate(total=Count("id"))
            .order_by("-total")
            .first()
        )

        if motivo_mas_frecuente:
            motivo_codigo = motivo_mas_frecuente[
                "motivo_ausencia"
            ]

            motivo_mas_frecuente_texto = dict(
                Asistencia.MotivoAusencia.choices
            ).get(
                motivo_codigo,
                motivo_codigo,
            )
        else:
            motivo_mas_frecuente_texto = "-"

        mes_porcentaje = (
            round(
                (
                    mes_presentes
                    / mes_total
                ) * 100,
                1,
            )
            if mes_total
            else 0
        )

        if mes_total == 0:
            estado_clase = "secondary"
            estado_texto = "Sin datos"

        elif mes_porcentaje >= 80:
            estado_clase = "success"
            estado_texto = "Buena asistencia"

        elif mes_porcentaje >= 50:
            estado_clase = "warning text-dark"
            estado_texto = "Asistencia media"

        else:
            estado_clase = "danger"
            estado_texto = "Baja asistencia"

        total_semana_presentes += semana_presentes
        total_semana_ausentes += semana_ausentes
        total_mes_presentes += mes_presentes
        total_mes_ausentes += mes_ausentes
        total_mes_pendientes += mes_pendientes

        datos.append({
            "jugador": jugador,

            "semana_total": semana_total,
            "semana_presentes": semana_presentes,
            "semana_tardes": semana_tardes,
            "semana_ausentes": semana_ausentes,
            "semana_pendientes": semana_pendientes,
            "semana_porcentaje": semana_porcentaje,

            "mes_total": mes_total,
            "mes_presentes": mes_presentes,
            "mes_tardes": mes_tardes,
            "mes_ausentes": mes_ausentes,
            "mes_pendientes": mes_pendientes,
            "mes_porcentaje": mes_porcentaje,

            "mes_ausencias_justificadas": (
                mes_ausencias_justificadas
            ),
            "mes_ausencias_sin_aviso": (
                mes_ausencias_sin_aviso
            ),
            "motivo_mas_frecuente": (
                motivo_mas_frecuente_texto
            ),

            "estado_clase": estado_clase,
            "estado_texto": estado_texto,
        })

    resumen_general = {
        "total_jugadores": total_jugadores,
        "total_semana_presentes": total_semana_presentes,
        "total_semana_ausentes": total_semana_ausentes,
        "total_mes_presentes": total_mes_presentes,
        "total_mes_ausentes": total_mes_ausentes,
        "total_mes_pendientes": total_mes_pendientes,
        "total_mes_asistencias": (
            total_mes_presentes
            + total_mes_ausentes
            + total_mes_pendientes
        ),
        "tiene_actividad_mes": (
            total_mes_presentes > 0
            or total_mes_ausentes > 0
            or total_mes_pendientes > 0
            or total_turnos_no_entrenados_mes > 0
        ),
        "dias_no_entrenados_semana": dias_no_entrenados_semana,
        "dias_no_entrenados_mes": dias_no_entrenados_mes,
        "total_turnos_no_entrenados_semana": (
            total_turnos_no_entrenados_semana
        ),
        "total_turnos_no_entrenados_mes": (
            total_turnos_no_entrenados_mes
        ),
    }

    return render(
        request,
        "asistencia/reportes.html",
        {
            "datos": datos,
            "hoy": hoy,
            "fecha_base": fecha_base,
            "inicio_semana": inicio_semana,
            "fin_semana": fin_semana,
            "inicio_mes": inicio_mes,
            "fin_mes": fin_mes,
            "semana_anterior": semana_anterior,
            "semana_siguiente": semana_siguiente,
            "mes_anterior": mes_anterior,
            "mes_siguiente": mes_siguiente,
            "resumen_general": resumen_general,
            "turnos_no_entrenados_semana": turnos_no_entrenados_semana,
            "turnos_no_entrenados_mes": turnos_no_entrenados_mes,
            "motivos_no_entrenamiento": motivos_no_entrenamiento,
        },
    )
    
    
@login_required
def resumen_dia(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        except ValueError:
            fecha = timezone.localdate()
    else:
        fecha = timezone.localdate()

    dia_anterior = fecha - timedelta(days=1)
    dia_siguiente = fecha + timedelta(days=1)

    entrenamientos = (
        Entrenamiento.objects
        .filter(
            club=club,
            fecha=fecha,
        )
        .select_related(
            "responsable_usuario",
            "turno_config",
        )
        .order_by(
            "turno_config__orden",
        )
    )

    turnos_resumen = []

    total_presentes = 0
    total_tardes = 0
    total_ausentes = 0
    total_trabajos = 0
    total_partidos = 0
    total_observaciones = 0
    total_ejercicios = 0

    def texto_campo(objeto, campos_posibles):
        for campo in campos_posibles:
            if hasattr(objeto, campo):
                valor = getattr(objeto, campo)
                if valor:
                    return valor
        return ""

    def tipo_trabajo(trabajo):
        if hasattr(trabajo, "get_tipo_display"):
            return trabajo.get_tipo_display()

        return texto_campo(
            trabajo,
            [
                "tipo",
                "tipo_trabajo",
            ],
        )

    def texto_observacion(observacion):
        return texto_campo(
            observacion,
            [
                "texto",
                "observacion",
                "detalle",
                "comentario",
            ],
        )

    def sets_del_partido(partido):
        if hasattr(partido, "sets"):
            return partido.sets.all()

        if hasattr(partido, "setpartido_set"):
            return partido.setpartido_set.all()

        return []

    def resultado_partido(partido):
        sets_j1 = 0
        sets_j2 = 0
        detalle_sets = []

        for set_partido in sets_del_partido(partido):
            puntos_j1 = texto_campo(
                set_partido,
                [
                    "puntos_jugador_1",
                    "puntos_j1",
                    "jugador_1_puntos",
                    "puntos_1",
                ],
            )

            puntos_j2 = texto_campo(
                set_partido,
                [
                    "puntos_jugador_2",
                    "puntos_j2",
                    "jugador_2_puntos",
                    "puntos_2",
                ],
            )

            if puntos_j1 == "" or puntos_j2 == "":
                continue

            try:
                puntos_j1_int = int(puntos_j1)
                puntos_j2_int = int(puntos_j2)
            except (TypeError, ValueError):
                continue

            detalle_sets.append(f"{puntos_j1_int}-{puntos_j2_int}")

            if puntos_j1_int > puntos_j2_int:
                sets_j1 += 1
            elif puntos_j2_int > puntos_j1_int:
                sets_j2 += 1

        if detalle_sets:
            return {
                "marcador": f"{sets_j1}-{sets_j2}",
                "sets": ", ".join(detalle_sets),
            }

        return {
            "marcador": "Sin sets cargados",
            "sets": "",
        }

    for entrenamiento in entrenamientos:
        asistencias = (
            Asistencia.objects
            .filter(entrenamiento=entrenamiento)
            .select_related("jugador")
            .order_by(
                "jugador__apellido",
                "jugador__nombre",
            )
        )

        presentes = asistencias.filter(estado="asistio").count()
        tardes = asistencias.filter(estado="tarde").count()
        ausentes = asistencias.filter(estado="ausente").count()

        ejercicios_turno = (
            EjercicioTurno.objects
            .filter(entrenamiento=entrenamiento)
            .select_related(
                "ejercicio",
                "ejercicio__categoria_config",
            )
            .order_by(
                "ejercicio__categoria_config__orden",
                "ejercicio__categoria_config__nombre",
                "ejercicio__nombre",
            )
        )

        ejercicios = []

        for ejercicio_turno in ejercicios_turno:
            ejercicio = ejercicio_turno.ejercicio

            categoria = ejercicio.nombre_categoria

            ejercicios.append(
                {
                    "nombre": ejercicio.nombre,
                    "categoria": categoria,
                }
            )

        trabajos_qs = (
            TrabajoTurno.objects
            .filter(entrenamiento=entrenamiento)
            .select_related("jugador_1", "jugador_2")
            .order_by("cambio", "id")
        )

        trabajos = []

        for trabajo in trabajos_qs:
            trabajos.append(
                {
                    "cambio": trabajo.cambio,
                    "tipo": tipo_trabajo(trabajo),
                    "jugador_1": getattr(trabajo, "jugador_1", None),
                    "jugador_2": getattr(trabajo, "jugador_2", None),
                    "detalle": texto_campo(
                        trabajo,
                        [
                            "detalle",
                            "descripcion",
                            "observacion",
                            "comentario",
                        ],
                    ),
                }
            )

        partidos_qs = (
            PartidoTurno.objects
            .filter(entrenamiento=entrenamiento)
            .select_related("jugador_1", "jugador_2")
            .prefetch_related("sets")
            .order_by("id")
        )

        partidos = []

        for partido in partidos_qs:
            resultado = resultado_partido(partido)

            partidos.append(
                {
                    "jugador_1": getattr(partido, "jugador_1", None),
                    "jugador_2": getattr(partido, "jugador_2", None),
                    "marcador": resultado["marcador"],
                    "sets": resultado["sets"],
                }
            )

        observaciones_qs = (
            ObservacionJugador.objects
            .filter(entrenamiento=entrenamiento)
            .select_related("jugador")
            .order_by("id")
        )

        observaciones = []

        for observacion in observaciones_qs:
            observaciones.append(
                {
                    "jugador": getattr(observacion, "jugador", None),
                    "texto": texto_observacion(observacion),
                }
            )

        total_presentes += presentes
        total_tardes += tardes
        total_ausentes += ausentes
        total_trabajos += len(trabajos)
        total_partidos += len(partidos)
        total_observaciones += len(observaciones)
        total_ejercicios += len(ejercicios)

        turnos_resumen.append(
            {
                "entrenamiento": entrenamiento,
                "entrenador_nombre": nombre_entrenador(entrenamiento),
                "presentes": presentes,
                "tardes": tardes,
                "ausentes": ausentes,
                "asistencias": asistencias,
                "ejercicios": ejercicios,
                "trabajos": trabajos,
                "partidos": partidos,
                "observaciones": observaciones,
            }
        )

    contexto = {
        "fecha": fecha,
        "dia_anterior": dia_anterior,
        "dia_siguiente": dia_siguiente,
        "turnos_resumen": turnos_resumen,
        "total_presentes": total_presentes,
        "total_tardes": total_tardes,
        "total_ausentes": total_ausentes,
        "total_trabajos": total_trabajos,
        "total_partidos": total_partidos,
        "total_observaciones": total_observaciones,
        "total_ejercicios": total_ejercicios,
    }

    return render(request, "asistencia/resumen_dia.html", contexto)
    


@requerir_club
def exportar_historial_jugador_excel(request, jugador_id):
    club = obtener_club_usuario(request.user)

    jugador = get_object_or_404(
        Jugador,
        id=jugador_id,
        club=club,
    )

    hoy = timezone.localdate()

    fecha_desde_str = request.GET.get("desde")
    fecha_hasta_str = request.GET.get("hasta")

    if fecha_desde_str:
        try:
            fecha_desde = datetime.strptime(
                fecha_desde_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_desde = hoy - timedelta(days=30)
    else:
        fecha_desde = hoy - timedelta(days=30)

    if fecha_hasta_str:
        try:
            fecha_hasta = datetime.strptime(
                fecha_hasta_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_hasta = hoy
    else:
        fecha_hasta = hoy

    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = (
            fecha_hasta,
            fecha_desde,
        )

    es_admin = PerfilUsuario.objects.filter(
        usuario=request.user,
        club=club,
        activo=True,
        rol=PerfilUsuario.Rol.ADMIN,
    ).exists()

    asistencias = (
        Asistencia.objects
        .filter(
            jugador=jugador,
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
            entrenamiento__no_se_entreno=False,
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
        )
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
        )
    )

    partidos = (
        PartidoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
        )
        .filter(
            Q(jugador_1=jugador)
            | Q(jugador_2=jugador)
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "jugador_1",
            "jugador_2",
        )
        .prefetch_related("sets")
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "-id",
        )
    )

    trabajos = (
        TrabajoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
        )
        .filter(
            Q(jugador_1=jugador)
            | Q(jugador_2=jugador)
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "jugador_1",
            "jugador_2",
        )
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "cambio",
        )
    )

    observaciones = (
        ObservacionJugador.objects
        .filter(
            jugador=jugador,
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                fecha_desde,
                fecha_hasta,
            ],
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "creada_por",
        )
        .order_by(
            "-entrenamiento__fecha",
            "-creada_el",
        )
    )

    entrenamientos_ids = list(
        asistencias.values_list(
            "entrenamiento_id",
            flat=True,
        )
    )

    ejercicios = (
        EjercicioTurno.objects
        .filter(
            entrenamiento_id__in=entrenamientos_ids,
            entrenamiento__club=club,
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "ejercicio",
            "ejercicio__categoria_config",
        )
        .order_by(
            "-entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "ejercicio__categoria_config__orden",
            "ejercicio__nombre",
        )
    )

    total_registros = asistencias.count()

    total_presentes = asistencias.filter(
        Q(estado="asistio")
        | Q(estado="tarde")
    ).count()

    total_tardes = asistencias.filter(
        estado="tarde",
    ).count()

    total_ausencias = asistencias.filter(
        estado="ausente",
    ).count()

    total_pendientes = asistencias.filter(
        estado="pendiente",
    ).count()

    porcentaje_asistencia = (
        round(
            (
                total_presentes
                / total_registros
            ) * 100,
            1,
        )
        if total_registros
        else 0
    )

    total_partidos = partidos.count()
    ganados = 0
    perdidos = 0
    empatados = 0

    for partido in partidos:
        sets_1 = partido.sets_jugador_1
        sets_2 = partido.sets_jugador_2

        if sets_1 == sets_2:
            empatados += 1
        elif partido.jugador_1_id == jugador.id:
            if sets_1 > sets_2:
                ganados += 1
            else:
                perdidos += 1
        else:
            if sets_2 > sets_1:
                ganados += 1
            else:
                perdidos += 1

    wb = Workbook()

    titulo_font = Font(
        bold=True,
        size=14,
        color="FFFFFF",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    seccion_font = Font(
        bold=True,
        color="FFFFFF",
    )

    titulo_fill = PatternFill(
        start_color="111827",
        end_color="111827",
        fill_type="solid",
    )

    header_fill = PatternFill(
        start_color="374151",
        end_color="374151",
        fill_type="solid",
    )

    seccion_fill = PatternFill(
        start_color="1F2937",
        end_color="1F2937",
        fill_type="solid",
    )

    center = Alignment(
        horizontal="center",
        vertical="center",
    )

    left = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    def turno_texto(entrenamiento):
        if entrenamiento.turno_config:
            texto = entrenamiento.turno_config.nombre

            if entrenamiento.turno_config.hora_inicio:
                texto += (
                    " · "
                    + entrenamiento.turno_config.hora_inicio.strftime(
                        "%H:%M"
                    )
                )

            if entrenamiento.turno_config.hora_fin:
                texto += (
                    "-"
                    + entrenamiento.turno_config.hora_fin.strftime(
                        "%H:%M"
                    )
                )

            return texto

        return entrenamiento.nombre_turno

    def aplicar_titulo(
        ws,
        texto,
        columnas,
    ):
        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=columnas,
        )

        celda = ws.cell(
            row=1,
            column=1,
            value=texto,
        )

        celda.font = titulo_font
        celda.fill = titulo_fill
        celda.alignment = center

    def aplicar_header(
        ws,
        fila,
    ):
        for celda in ws[fila]:
            celda.font = header_font
            celda.fill = header_fill
            celda.alignment = center

    def aplicar_seccion(
        ws,
        fila,
        texto,
        columnas,
    ):
        ws.merge_cells(
            start_row=fila,
            start_column=1,
            end_row=fila,
            end_column=columnas,
        )

        celda = ws.cell(
            row=fila,
            column=1,
            value=texto,
        )

        celda.font = seccion_font
        celda.fill = seccion_fill
        celda.alignment = left

    def ajustar_columnas(ws):
        for fila in ws.iter_rows():
            for celda in fila:
                if celda.row != 1:
                    celda.alignment = left

        for numero_columna in range(
            1,
            ws.max_column + 1,
        ):
            letra = get_column_letter(
                numero_columna
            )

            ancho_maximo = 0

            for numero_fila in range(
                1,
                ws.max_row + 1,
            ):
                valor = ws.cell(
                    row=numero_fila,
                    column=numero_columna,
                ).value

                if valor is not None:
                    ancho_maximo = max(
                        ancho_maximo,
                        len(str(valor)),
                    )

            ws.column_dimensions[
                letra
            ].width = min(
                ancho_maximo + 3,
                45,
            )

    nombre_jugador = str(jugador)

    titulo_base = (
        f"HISTORIAL - {nombre_jugador}"
    )

    # =========================
    # HOJA 1: RESUMEN
    # =========================
    ws = wb.active
    ws.title = "Resumen"

    aplicar_titulo(
        ws,
        titulo_base,
        4,
    )

    aplicar_seccion(
        ws,
        3,
        "DATOS DEL JUGADOR",
        4,
    )

    ws.append([
        "Campo",
        "Valor",
        "",
        "",
    ])
    aplicar_header(ws, 4)

    ws.append([
        "Jugador",
        nombre_jugador,
        "",
        "",
    ])

    ws.append([
        "Categoría",
        (
            jugador.categoria.nombre
            if jugador.categoria
            else "Sin categoría"
        ),
        "",
        "",
    ])

    ws.append([
        "Estado",
        (
            "Activo"
            if jugador.activo
            else "Inactivo"
        ),
        "",
        "",
    ])

    ws.append([
        "Fecha de alta",
        jugador.fecha_alta.strftime("%d/%m/%Y"),
        "",
        "",
    ])

    ws.append([
        "Fecha de baja",
        (
            jugador.fecha_baja.strftime("%d/%m/%Y")
            if jugador.fecha_baja
            else "-"
        ),
        "",
        "",
    ])

    ws.append([
        "Período exportado",
        (
            f"{fecha_desde.strftime('%d/%m/%Y')} "
            f"al {fecha_hasta.strftime('%d/%m/%Y')}"
        ),
        "",
        "",
    ])

    aplicar_seccion(
        ws,
        12,
        "RESUMEN DEL PERÍODO",
        4,
    )

    ws.append([
        "Indicador",
        "Valor",
        "",
        "",
    ])
    aplicar_header(ws, 13)

    resumen_filas = [
        ["Registros de asistencia", total_registros],
        ["Presentes", total_presentes],
        ["Tardes", total_tardes],
        ["Ausencias", total_ausencias],
        ["Pendientes", total_pendientes],
        ["% asistencia", porcentaje_asistencia / 100],
        ["Partidos", total_partidos],
        ["Partidos ganados", ganados],
        ["Partidos perdidos", perdidos],
        ["Partidos empatados", empatados],
        ["Trabajos", trabajos.count()],
        ["Observaciones", observaciones.count()],
    ]

    for etiqueta, valor in resumen_filas:
        ws.append([
            etiqueta,
            valor,
            "",
            "",
        ])

    ws.cell(
        row=19,
        column=2,
    ).number_format = "0.0%"

    ajustar_columnas(ws)

    # =========================
    # HOJA 2: ASISTENCIAS
    # =========================
    ws = wb.create_sheet("Asistencias")

    aplicar_titulo(
        ws,
        titulo_base,
        6,
    )

    ws.append([])
    ws.append([
        "Fecha",
        "Turno",
        "Estado",
        "Motivo ausencia",
        "Detalle ausencia",
        "Entrenador",
    ])
    aplicar_header(ws, 3)

    for asistencia in asistencias:
        entrenamiento = asistencia.entrenamiento

        if asistencia.estado == "asistio":
            estado = "Asistió"
        elif asistencia.estado == "tarde":
            estado = "Tarde"
        elif asistencia.estado == "ausente":
            estado = "Ausente"
        else:
            estado = "Pendiente"

        motivo = (
            asistencia.get_motivo_ausencia_display()
            if asistencia.motivo_ausencia
            else "-"
        )

        ws.append([
            entrenamiento.fecha,
            turno_texto(entrenamiento),
            estado,
            motivo,
            asistencia.detalle_ausencia or "-",
            entrenamiento.nombre_entrenador,
        ])

        ws.cell(
            row=ws.max_row,
            column=1,
        ).number_format = "DD/MM/YYYY"

    ws.freeze_panes = "A4"

    if ws.max_row >= 4:
        ws.auto_filter.ref = (
            f"A3:F{ws.max_row}"
        )

    ajustar_columnas(ws)

    # =========================
    # HOJA 3: PARTIDOS
    # =========================
    ws = wb.create_sheet("Partidos")

    aplicar_titulo(
        ws,
        titulo_base,
        7,
    )

    ws.append([])
    ws.append([
        "Fecha",
        "Turno",
        "Rival",
        "Resultado",
        "Condición",
        "Sets",
        "Detalle",
    ])
    aplicar_header(ws, 3)

    for partido in partidos:
        sets_1 = partido.sets_jugador_1
        sets_2 = partido.sets_jugador_2

        if partido.jugador_1_id == jugador.id:
            rival = partido.jugador_2
            sets_jugador = sets_1
            sets_rival = sets_2
        else:
            rival = partido.jugador_1
            sets_jugador = sets_2
            sets_rival = sets_1

        if sets_jugador > sets_rival:
            condicion = "Ganado"
        elif sets_jugador < sets_rival:
            condicion = "Perdido"
        else:
            condicion = "Empatado"

        sets_texto = []

        for set_partido in partido.sets.all():
            if partido.jugador_1_id == jugador.id:
                puntos_jugador = (
                    set_partido.puntos_jugador_1
                )
                puntos_rival = (
                    set_partido.puntos_jugador_2
                )
            else:
                puntos_jugador = (
                    set_partido.puntos_jugador_2
                )
                puntos_rival = (
                    set_partido.puntos_jugador_1
                )

            sets_texto.append(
                f"{puntos_jugador}-{puntos_rival}"
            )

        ws.append([
            partido.entrenamiento.fecha,
            turno_texto(
                partido.entrenamiento
            ),
            str(rival),
            f"{sets_jugador}-{sets_rival}",
            condicion,
            (
                " / ".join(sets_texto)
                if sets_texto
                else "-"
            ),
            partido.detalle or "-",
        ])

        ws.cell(
            row=ws.max_row,
            column=1,
        ).number_format = "DD/MM/YYYY"

    ws.freeze_panes = "A4"

    if ws.max_row >= 4:
        ws.auto_filter.ref = (
            f"A3:G{ws.max_row}"
        )

    ajustar_columnas(ws)

    # =========================
    # HOJA 4: TRABAJOS
    # =========================
    ws = wb.create_sheet("Trabajos")

    aplicar_titulo(
        ws,
        titulo_base,
        7,
    )

    ws.append([])
    ws.append([
        "Fecha",
        "Turno",
        "Cambio",
        "Tipo",
        "Compañero",
        "Detalle",
        "Jugador",
    ])
    aplicar_header(ws, 3)

    for trabajo in trabajos:
        if trabajo.jugador_1_id == jugador.id:
            companero = trabajo.jugador_2
        else:
            companero = trabajo.jugador_1

        ws.append([
            trabajo.entrenamiento.fecha,
            turno_texto(
                trabajo.entrenamiento
            ),
            trabajo.cambio,
            trabajo.get_tipo_display(),
            (
                str(companero)
                if companero
                else "-"
            ),
            trabajo.detalle or "-",
            nombre_jugador,
        ])

        ws.cell(
            row=ws.max_row,
            column=1,
        ).number_format = "DD/MM/YYYY"

    ws.freeze_panes = "A4"

    if ws.max_row >= 4:
        ws.auto_filter.ref = (
            f"A3:G{ws.max_row}"
        )

    ajustar_columnas(ws)

    # =========================
    # HOJA 5: OBSERVACIONES
    # =========================
    ws = wb.create_sheet("Observaciones")

    aplicar_titulo(
        ws,
        titulo_base,
        5,
    )

    ws.append([])
    ws.append([
        "Fecha",
        "Turno",
        "Observación",
        "Creada por",
        "Fecha de carga",
    ])
    aplicar_header(ws, 3)

    for observacion in observaciones:
        ws.append([
            observacion.entrenamiento.fecha,
            turno_texto(
                observacion.entrenamiento
            ),
            observacion.texto,
            (
                observacion.creada_por.get_full_name()
                or observacion.creada_por.username
                if observacion.creada_por
                else "-"
            ),
            observacion.creada_el.replace(
                tzinfo=None
            ),
        ])

        ws.cell(
            row=ws.max_row,
            column=1,
        ).number_format = "DD/MM/YYYY"

        ws.cell(
            row=ws.max_row,
            column=5,
        ).number_format = "DD/MM/YYYY HH:MM"

    ws.freeze_panes = "A4"

    if ws.max_row >= 4:
        ws.auto_filter.ref = (
            f"A3:E{ws.max_row}"
        )

    ajustar_columnas(ws)

    # =========================
    # HOJA 6: EJERCICIOS
    # =========================
    ws = wb.create_sheet("Ejercicios")

    aplicar_titulo(
        ws,
        titulo_base,
        4,
    )

    ws.append([])
    ws.append([
        "Fecha",
        "Turno",
        "Categoría",
        "Ejercicio",
    ])
    aplicar_header(ws, 3)

    for item in ejercicios:
        ws.append([
            item.entrenamiento.fecha,
            turno_texto(
                item.entrenamiento
            ),
            (
                item.ejercicio.categoria_config.nombre
                if item.ejercicio.categoria_config
                else "Sin categoría"
            ),
            item.ejercicio.nombre,
        ])

        ws.cell(
            row=ws.max_row,
            column=1,
        ).number_format = "DD/MM/YYYY"

    ws.freeze_panes = "A4"

    if ws.max_row >= 4:
        ws.auto_filter.ref = (
            f"A3:D{ws.max_row}"
        )

    ajustar_columnas(ws)

    # =========================
    # HOJA 7: CUOTAS
    # SOLO ADMIN
    # =========================
    if es_admin:
        ws = wb.create_sheet("Cuotas")

        aplicar_titulo(
            ws,
            titulo_base,
            5,
        )

        ws.append([])
        ws.append([
            "Período",
            "Estado",
            "Fecha de pago",
            "Observación",
            "Jugador",
        ])
        aplicar_header(ws, 3)

        fecha_inicio = jugador.fecha_alta.replace(
            day=1
        )

        fecha_fin_membresia = (
            jugador.fecha_baja
            if jugador.fecha_baja
            else hoy
        )

        fecha_fin = min(
            fecha_fin_membresia,
            hoy,
        ).replace(day=1)

        pagos = {
            (pago.anio, pago.mes): pago
            for pago in PagoJugador.objects.filter(
                jugador=jugador,
            )
        }

        cursor = fecha_inicio

        while cursor <= fecha_fin:
            pago = pagos.get(
                (
                    cursor.year,
                    cursor.month,
                )
            )

            pagado = (
                pago.pagado
                if pago
                else False
            )

            ws.append([
                cursor,
                (
                    "Pagado"
                    if pagado
                    else "Pendiente"
                ),
                (
                    pago.fecha_pago
                    if (
                        pago
                        and pago.fecha_pago
                    )
                    else "-"
                ),
                (
                    pago.observacion
                    if (
                        pago
                        and pago.observacion
                    )
                    else "-"
                ),
                nombre_jugador,
            ])

            ws.cell(
                row=ws.max_row,
                column=1,
            ).number_format = "MM/YYYY"

            if (
                pago
                and pago.fecha_pago
            ):
                ws.cell(
                    row=ws.max_row,
                    column=3,
                ).number_format = "DD/MM/YYYY"

            if cursor.month == 12:
                cursor = cursor.replace(
                    year=cursor.year + 1,
                    month=1,
                )
            else:
                cursor = cursor.replace(
                    month=cursor.month + 1,
                )

        ws.freeze_panes = "A4"

        if ws.max_row >= 4:
            ws.auto_filter.ref = (
                f"A3:E{ws.max_row}"
            )

        ajustar_columnas(ws)

    response = HttpResponse(
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    nombre_archivo = (
        f"historial_jugador_{jugador.id}_"
        f"{fecha_desde.strftime('%Y%m%d')}_"
        f"{fecha_hasta.strftime('%Y%m%d')}.xlsx"
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{nombre_archivo}"'
    )

    wb.save(response)

    return response


@requerir_club
def exportar_asistencia_excel(request):
    club = obtener_club_usuario(request.user)
    hoy = timezone.localdate()

    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha_base = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_base = hoy
    else:
        fecha_base = hoy

    inicio_mes = fecha_base.replace(day=1)

    if inicio_mes.month == 12:
        inicio_mes_siguiente = inicio_mes.replace(
            year=inicio_mes.year + 1,
            month=1,
        )
    else:
        inicio_mes_siguiente = inicio_mes.replace(
            month=inicio_mes.month + 1,
        )

    fin_mes = inicio_mes_siguiente - timedelta(days=1)

    asistencias = (
        Asistencia.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
            entrenamiento__no_se_entreno=False,
        )
        .select_related(
            "jugador",
            "jugador__categoria",
            "entrenamiento",
            "entrenamiento__turno_config",
            "entrenamiento__responsable_usuario",
        )
        .order_by(
            F("jugador__categoria__orden").asc(
                nulls_last=True
            ),
            F("jugador__categoria__nombre").asc(
                nulls_last=True
            ),
            "jugador__apellido",
            "jugador__nombre",
            "entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
        )
    )

    jugadores = (
        Jugador.objects
        .filter(
            club=club,
            fecha_alta__lte=fin_mes,
        )
        .filter(
            Q(fecha_baja__isnull=True)
            | Q(fecha_baja__gte=inicio_mes)
        )
        .select_related("categoria")
        .order_by(
            F("categoria__orden").asc(
                nulls_last=True
            ),
            F("categoria__nombre").asc(
                nulls_last=True
            ),
            "apellido",
            "nombre",
        )
    )

    workbook = Workbook()

    hoja_resumen = workbook.active
    hoja_resumen.title = "Resumen"

    titulo_font = Font(
        bold=True,
        size=14,
        color="FFFFFF",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    titulo_fill = PatternFill(
        start_color="111827",
        end_color="111827",
        fill_type="solid",
    )

    header_fill = PatternFill(
        start_color="374151",
        end_color="374151",
        fill_type="solid",
    )

    center = Alignment(
        horizontal="center",
        vertical="center",
    )

    left = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    def aplicar_titulo(
        hoja,
        texto,
        columnas,
    ):
        hoja.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=columnas,
        )

        celda = hoja.cell(
            row=1,
            column=1,
            value=texto,
        )

        celda.font = titulo_font
        celda.fill = titulo_fill
        celda.alignment = center

    def aplicar_encabezados(
        hoja,
        fila,
    ):
        for celda in hoja[fila]:
            celda.font = header_font
            celda.fill = header_fill
            celda.alignment = center

    def ajustar_columnas(hoja):
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.row != 1:
                    celda.alignment = left

        for numero_columna in range(
            1,
            hoja.max_column + 1,
        ):
            letra = get_column_letter(
                numero_columna
            )

            ancho_maximo = 0

            for numero_fila in range(
                1,
                hoja.max_row + 1,
            ):
                valor = hoja.cell(
                    row=numero_fila,
                    column=numero_columna,
                ).value

                if valor is not None:
                    ancho_maximo = max(
                        ancho_maximo,
                        len(str(valor)),
                    )

            hoja.column_dimensions[
                letra
            ].width = min(
                ancho_maximo + 3,
                40,
            )

    nombre_club = getattr(
        club,
        "nombre",
        str(club),
    )

    titulo = (
        f"ASISTENCIA - {nombre_club} - "
        f"{inicio_mes.strftime('%m/%Y')}"
    )

    aplicar_titulo(
        hoja_resumen,
        titulo,
        8,
    )

    hoja_resumen.append([])
    hoja_resumen.append([
        "Jugador",
        "Categoría",
        "Registros",
        "Presentes",
        "Tardes",
        "Ausencias",
        "Pendientes",
        "% asistencia",
    ])

    aplicar_encabezados(
        hoja_resumen,
        3,
    )

    for jugador in jugadores:
        registros = asistencias.filter(
            jugador=jugador,
        )

        total = registros.count()

        presentes = registros.filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        ).count()

        tardes = registros.filter(
            estado="tarde",
        ).count()

        ausencias = registros.filter(
            estado="ausente",
        ).count()

        pendientes = registros.filter(
            estado="pendiente",
        ).count()

        porcentaje = (
            round(
                (
                    presentes
                    / total
                ) * 100,
                1,
            )
            if total
            else 0
        )

        hoja_resumen.append([
            str(jugador),
            (
                jugador.categoria.nombre
                if jugador.categoria
                else "Sin categoría"
            ),
            total,
            presentes,
            tardes,
            ausencias,
            pendientes,
            porcentaje / 100,
        ])

        hoja_resumen.cell(
            row=hoja_resumen.max_row,
            column=8,
        ).number_format = "0.0%"

    hoja_resumen.freeze_panes = "A4"

    if hoja_resumen.max_row >= 4:
        hoja_resumen.auto_filter.ref = (
            f"A3:H{hoja_resumen.max_row}"
        )

    ajustar_columnas(
        hoja_resumen
    )

    hoja_detalle = workbook.create_sheet(
        "Asistencia detallada"
    )

    aplicar_titulo(
        hoja_detalle,
        titulo,
        8,
    )

    hoja_detalle.append([])
    hoja_detalle.append([
        "Fecha",
        "Turno",
        "Jugador",
        "Categoría",
        "Estado",
        "Motivo ausencia",
        "Detalle ausencia",
        "Entrenador",
    ])

    aplicar_encabezados(
        hoja_detalle,
        3,
    )

    for asistencia in asistencias:
        entrenamiento = asistencia.entrenamiento

        if entrenamiento.turno_config:
            turno = entrenamiento.turno_config.nombre

            if entrenamiento.turno_config.hora_inicio:
                turno += (
                    " · "
                    + entrenamiento.turno_config.hora_inicio.strftime(
                        "%H:%M"
                    )
                )

            if entrenamiento.turno_config.hora_fin:
                turno += (
                    "-"
                    + entrenamiento.turno_config.hora_fin.strftime(
                        "%H:%M"
                    )
                )
        else:
            turno = entrenamiento.nombre_turno

        if asistencia.estado == "asistio":
            estado = "Asistió"
        elif asistencia.estado == "tarde":
            estado = "Tarde"
        elif asistencia.estado == "ausente":
            estado = "Ausente"
        else:
            estado = "Pendiente"

        if (
            asistencia.estado == "ausente"
            and asistencia.motivo_ausencia
        ):
            motivo = (
                asistencia.get_motivo_ausencia_display()
            )
        else:
            motivo = "-"

        detalle = (
            asistencia.detalle_ausencia
            if asistencia.detalle_ausencia
            else "-"
        )

        entrenador = (
            entrenamiento.nombre_entrenador
        )

        hoja_detalle.append([
            entrenamiento.fecha,
            turno,
            str(asistencia.jugador),
            (
                asistencia.jugador.categoria.nombre
                if asistencia.jugador.categoria
                else "Sin categoría"
            ),
            estado,
            motivo,
            detalle,
            entrenador,
        ])

        hoja_detalle.cell(
            row=hoja_detalle.max_row,
            column=1,
        ).number_format = "DD/MM/YYYY"

    hoja_detalle.freeze_panes = "A4"

    if hoja_detalle.max_row >= 4:
        hoja_detalle.auto_filter.ref = (
            f"A3:H{hoja_detalle.max_row}"
        )

    ajustar_columnas(
        hoja_detalle
    )

    response = HttpResponse(
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    nombre_archivo = (
        "asistencia_"
        f"{inicio_mes.year}_"
        f"{inicio_mes.month:02d}.xlsx"
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{nombre_archivo}"'
    )

    workbook.save(
        response
    )

    return response


@login_required
def exportar_reporte_mensual(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    hoy = timezone.localdate()
    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha_base = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_base = hoy
    else:
        fecha_base = hoy

    inicio_mes = fecha_base.replace(day=1)

    if inicio_mes.month == 12:
        inicio_mes_siguiente = inicio_mes.replace(
            year=inicio_mes.year + 1,
            month=1,
        )
    else:
        inicio_mes_siguiente = inicio_mes.replace(
            month=inicio_mes.month + 1,
        )

    fin_mes = inicio_mes_siguiente - timedelta(days=1)

    wb = Workbook()

    # =========================
    # ESTILOS
    # =========================
    titulo_font = Font(bold=True, size=14, color="FFFFFF")
    seccion_font = Font(bold=True, size=12, color="FFFFFF")
    header_font = Font(bold=True, color="FFFFFF")
    normal_bold = Font(bold=True)

    titulo_fill = PatternFill(
        start_color="111827",
        end_color="111827",
        fill_type="solid",
    )

    seccion_fill = PatternFill(
        start_color="1F2937",
        end_color="1F2937",
        fill_type="solid",
    )

    header_fill = PatternFill(
        start_color="374151",
        end_color="374151",
        fill_type="solid",
    )

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def ajustar_columnas(ws):
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = left

        for column_index in range(1, ws.max_column + 1):
            max_length = 0
            column_letter = get_column_letter(column_index)

            for row_index in range(1, ws.max_row + 1):
                cell = ws.cell(
                    row=row_index,
                    column=column_index,
                )

                value = cell.value

                if value is not None:
                    max_length = max(
                        max_length,
                        len(str(value)),
                    )

            ws.column_dimensions[column_letter].width = min(
                max_length + 3,
                38,
            )

    def aplicar_header(ws, row_number):
        for cell in ws[row_number]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

    def aplicar_titulo(ws, row_number, texto):
        ws.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=8,
        )

        cell = ws.cell(row=row_number, column=1)
        cell.value = texto
        cell.font = titulo_font
        cell.fill = titulo_fill
        cell.alignment = center

    def aplicar_seccion(ws, row_number, texto):
        ws.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=8,
        )

        cell = ws.cell(row=row_number, column=1)
        cell.value = texto
        cell.font = seccion_font
        cell.fill = seccion_fill
        cell.alignment = left

    # =========================
    # DATOS BASE
    # =========================
    asistencias_mes = Asistencia.objects.filter(
        entrenamiento__club=club,
        jugador__club=club,
        entrenamiento__fecha__range=[
            inicio_mes,
            fin_mes,
        ],
        entrenamiento__no_se_entreno=False,
    ).select_related(
        "jugador",
        "entrenamiento",
        "entrenamiento__responsable_usuario",
        "entrenamiento__turno_config",
    )

    turnos_no_entrenados = Entrenamiento.objects.filter(
        club=club,
        fecha__range=[
            inicio_mes,
            fin_mes,
        ],
        no_se_entreno=True,
    ).select_related(
        "turno_config",
    ).order_by(
        "fecha",
        "turno_config__orden",
    )

    partidos_mes = (
        PartidoTurno.objects
        .filter(
            entrenamiento__club=club,
            jugador_1__club=club,
            jugador_2__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
        )
        .select_related(
            "entrenamiento",
            "entrenamiento__turno_config",
            "jugador_1",
            "jugador_2",
        )
        .prefetch_related("sets")
        .order_by(
            "entrenamiento__fecha",
            "entrenamiento__turno_config__orden",
            "id",
        )
    )

    jugadores = Jugador.objects.filter(
        club=club,
    ).order_by(
        "apellido",
        "nombre",
    )

    jugadores_activos = jugadores.filter(
        activo=True,
    ).count()

    jugadores_totales = jugadores.count()

    total_presentes = asistencias_mes.filter(
        Q(estado="asistio")
        | Q(estado="tarde")
    ).count()

    total_tardes = asistencias_mes.filter(
        estado="tarde",
    ).count()

    total_ausencias = asistencias_mes.filter(
        estado="ausente",
    ).count()

    total_pendientes = asistencias_mes.filter(
        estado="pendiente",
    ).count()

    # =========================
    # HOJA 1: PANEL GENERAL
    # =========================
    ws = wb.active
    ws.title = "Panel general"

    fila = 1

    aplicar_titulo(
        ws,
        fila,
        f"REPORTE MENSUAL - {club.nombre} - {inicio_mes.strftime('%m/%Y')}",
    )

    fila += 2

    aplicar_seccion(ws, fila, "RESUMEN DEL MES")
    fila += 1

    ws.append(["Campo", "Valor"])
    aplicar_header(ws, fila)
    fila += 1

    resumen_filas = [
        ["Mes", inicio_mes.strftime("%m/%Y")],
        ["Desde", inicio_mes.strftime("%d/%m/%Y")],
        ["Hasta", fin_mes.strftime("%d/%m/%Y")],
        ["Jugadores totales", jugadores_totales],
        ["Jugadores activos", jugadores_activos],
        ["Presentes", total_presentes],
        ["Tardes", total_tardes],
        ["Ausencias", total_ausencias],
        ["Pendientes", total_pendientes],
        ["Turnos sin entrenamiento", turnos_no_entrenados.count()],
        ["Partidos cargados", partidos_mes.count()],
    ]

    for item in resumen_filas:
        ws.append(item)
        fila += 1

    fila += 2

    aplicar_seccion(ws, fila, "JUGADORES")
    fila += 1

    ws.append([
        "ID",
        "Nombre",
        "Apellido",
        "Jugador completo",
        "Activo",
    ])

    aplicar_header(ws, fila)
    fila += 1

    for jugador in jugadores:
        ws.append([
            jugador.id,
            jugador.nombre,
            jugador.apellido,
            str(jugador),
            "Sí" if jugador.activo else "No",
        ])

        fila += 1

    fila += 2

    aplicar_seccion(ws, fila, "ASISTENCIA POR JUGADOR")
    fila += 1

    ws.append([
        "Jugador",
        "Turnos cargados",
        "Presentes",
        "Tardes",
        "Ausencias",
        "Pendientes",
        "% asistencia",
        "Motivo frecuente",
    ])

    aplicar_header(ws, fila)
    fila += 1

    datos_asistencia_jugadores = []

    for jugador in jugadores.filter(activo=True):
        asistencias_jugador = asistencias_mes.filter(
            jugador=jugador,
        )

        total = asistencias_jugador.count()

        presentes = asistencias_jugador.filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        ).count()

        tardes = asistencias_jugador.filter(
            estado="tarde",
        ).count()

        ausencias = asistencias_jugador.filter(
            estado="ausente",
        )

        ausentes = ausencias.count()

        pendientes = asistencias_jugador.filter(
            estado="pendiente",
        ).count()

        porcentaje = (
            round(
                (
                    presentes
                    / total
                ) * 100,
                1,
            )
            if total
            else 0
        )

        justificadas = (
            ausencias
            .exclude(motivo_ausencia="")
            .exclude(motivo_ausencia="sin_aviso")
            .count()
        )

        sin_aviso = ausencias.filter(
            Q(motivo_ausencia="")
            | Q(motivo_ausencia="sin_aviso")
        ).count()

        motivo_mas_frecuente = (
            ausencias
            .exclude(motivo_ausencia="")
            .values("motivo_ausencia")
            .annotate(total=Count("id"))
            .order_by("-total")
            .first()
        )

        if motivo_mas_frecuente:
            motivo_codigo = motivo_mas_frecuente["motivo_ausencia"]

            motivo_texto = dict(
                Asistencia.MotivoAusencia.choices
            ).get(
                motivo_codigo,
                motivo_codigo,
            )
        else:
            motivo_texto = "-"

        datos_asistencia_jugadores.append({
            "jugador": jugador,
            "total": total,
            "presentes": presentes,
            "tardes": tardes,
            "ausentes": ausentes,
            "pendientes": pendientes,
            "porcentaje": porcentaje,
            "justificadas": justificadas,
            "sin_aviso": sin_aviso,
            "motivo_texto": motivo_texto,
        })

        ws.append([
            str(jugador),
            total,
            presentes,
            tardes,
            ausentes,
            pendientes,
            porcentaje,
            motivo_texto,
        ])

        fila += 1

    fila += 2

    aplicar_seccion(ws, fila, "TURNOS SIN ENTRENAMIENTO")
    fila += 1

    ws.append([
        "Fecha",
        "Turno",
        "Motivo",
        "Detalle",
    ])

    aplicar_header(ws, fila)
    fila += 1

    if turnos_no_entrenados.exists():
        for entrenamiento in turnos_no_entrenados:
            ws.append([
                entrenamiento.fecha.strftime("%d/%m/%Y"),
                entrenamiento.nombre_turno_completo,
                entrenamiento.get_motivo_no_entrenamiento_display(),
                entrenamiento.detalle_no_entrenamiento or "-",
            ])

            fila += 1
    else:
        ws.append([
            "-",
            "-",
            "No hubo turnos sin entrenamiento",
            "-",
        ])

        fila += 1

    ajustar_columnas(ws)

    # =========================
    # HOJA 2: ASISTENCIA DETALLADA
    # =========================
    ws = wb.create_sheet("Asistencia detallada")

    ws.append([
        "Fecha",
        "Turno",
        "Jugador",
        "Estado",
        "Motivo ausencia",
        "Detalle ausencia",
        "Entrenador",
    ])

    aplicar_header(ws, 1)

    asistencias_detalladas = asistencias_mes.order_by(
        "entrenamiento__fecha",
        "entrenamiento__turno_config__orden",
        "jugador__apellido",
        "jugador__nombre",
    )

    for asistencia in asistencias_detalladas:
        entrenamiento = asistencia.entrenamiento

        entrenador = nombre_entrenador(entrenamiento)

        ws.append([
            entrenamiento.fecha.strftime("%d/%m/%Y"),
            entrenamiento.nombre_turno_completo,
            str(asistencia.jugador),
            asistencia.get_estado_display(),
            asistencia.get_motivo_ausencia_display()
            if asistencia.motivo_ausencia
            else "-",
            asistencia.detalle_ausencia or "-",
            entrenador,
        ])

    ajustar_columnas(ws)

    # =========================
    # HOJA 3: AUSENCIAS
    # =========================
    ws = wb.create_sheet("Ausencias")

    ws.append([
        "Fecha",
        "Turno",
        "Jugador",
        "Motivo",
        "Detalle",
        "Entrenador",
    ])

    aplicar_header(ws, 1)

    ausencias_mes = asistencias_mes.filter(
        estado="ausente",
    ).order_by(
        "entrenamiento__fecha",
        "entrenamiento__turno_config__orden",
        "jugador__apellido",
        "jugador__nombre",
    )

    for asistencia in ausencias_mes:
        entrenamiento = asistencia.entrenamiento

        entrenador = nombre_entrenador(entrenamiento)

        ws.append([
            entrenamiento.fecha.strftime("%d/%m/%Y"),
            entrenamiento.nombre_turno_completo,
            str(asistencia.jugador),
            asistencia.get_motivo_ausencia_display()
            if asistencia.motivo_ausencia
            else "Sin aviso",
            asistencia.detalle_ausencia or "-",
            entrenador,
        ])

    ajustar_columnas(ws)

    # =========================
    # HOJA 4: PARTIDOS
    # =========================
    ws = wb.create_sheet("Partidos")

    ws.append([
        "Fecha",
        "Turno",
        "Jugador 1",
        "Jugador 2",
        "Resultado",
        "Ganador",
        "Sets",
        "Detalle",
    ])

    aplicar_header(ws, 1)

    for partido in partidos_mes:
        sets_jugador_1 = partido.sets_jugador_1
        sets_jugador_2 = partido.sets_jugador_2

        if sets_jugador_1 > sets_jugador_2:
            ganador = str(partido.jugador_1)
        elif sets_jugador_2 > sets_jugador_1:
            ganador = str(partido.jugador_2)
        else:
            ganador = "-"

        sets_texto = []

        for set_partido in partido.sets.all():
            sets_texto.append(
                f"{set_partido.puntos_jugador_1}-{set_partido.puntos_jugador_2}"
            )

        ws.append([
            partido.entrenamiento.fecha.strftime("%d/%m/%Y"),
            partido.entrenamiento.nombre_turno_completo,
            str(partido.jugador_1),
            str(partido.jugador_2),
            f"{sets_jugador_1}-{sets_jugador_2}",
            ganador,
            " / ".join(sets_texto) if sets_texto else "-",
            partido.detalle or "-",
        ])

    ajustar_columnas(ws)

    # =========================
    # RESPUESTA DESCARGA
    # =========================
    nombre_archivo = (
        f"reporte_mensual_"
        f"{inicio_mes.strftime('%Y_%m')}.xlsx"
    )

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{nombre_archivo}"'
    )

    wb.save(response)

    return response

@login_required
def dashboard_mensual(request):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    hoy = timezone.localdate()
    fecha_str = request.GET.get("fecha")

    if fecha_str:
        try:
            fecha_base = datetime.strptime(
                fecha_str,
                "%Y-%m-%d",
            ).date()
        except ValueError:
            fecha_base = hoy
    else:
        fecha_base = hoy

    inicio_mes = fecha_base.replace(day=1)

    if inicio_mes.month == 12:
        inicio_mes_siguiente = inicio_mes.replace(
            year=inicio_mes.year + 1,
            month=1,
        )
    else:
        inicio_mes_siguiente = inicio_mes.replace(
            month=inicio_mes.month + 1,
        )

    fin_mes = inicio_mes_siguiente - timedelta(days=1)

    if fin_mes > hoy:
        fin_mes = hoy

    mes_anterior = inicio_mes - timedelta(days=1)
    mes_siguiente = inicio_mes_siguiente

    asistencias_mes = (
        Asistencia.objects
        .filter(
            entrenamiento__club=club,
            jugador__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
            entrenamiento__no_se_entreno=False,
        )
        .select_related(
            "jugador",
            "entrenamiento",
        )
    )

    entrenamientos_mes = Entrenamiento.objects.filter(
        club=club,
        fecha__range=[
            inicio_mes,
            fin_mes,
        ],
    )

    turnos_no_entrenados = (
        entrenamientos_mes
        .filter(no_se_entreno=True)
        .select_related("turno_config")
        .order_by(
            "-fecha",
            "turno_config__orden",
        )
    )

    motivos_no_entrenamiento = (
        turnos_no_entrenados
        .values("motivo_no_entrenamiento")
        .annotate(total=Count("id"))
        .order_by("-total")
    )

    opciones_motivos = dict(
        Entrenamiento.MotivoNoEntrenamiento.choices
    )

    motivos = []

    for item in motivos_no_entrenamiento:
        codigo = item["motivo_no_entrenamiento"]

        motivos.append({
            "motivo": opciones_motivos.get(
                codigo,
                codigo or "Sin motivo",
            ),
            "total": item["total"],
        })

    ids_con_asistencia = set(
        Asistencia.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
        )
        .values_list(
            "entrenamiento_id",
            flat=True,
        )
    )

    ids_con_trabajos = set(
        TrabajoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
        )
        .values_list(
            "entrenamiento_id",
            flat=True,
        )
    )

    ids_con_partidos = set(
        PartidoTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
        )
        .values_list(
            "entrenamiento_id",
            flat=True,
        )
    )

    ids_con_ejercicios = set(
        EjercicioTurno.objects
        .filter(
            entrenamiento__club=club,
            entrenamiento__fecha__range=[
                inicio_mes,
                fin_mes,
            ],
        )
        .values_list(
            "entrenamiento_id",
            flat=True,
        )
    )

    ids_turnos_con_actividad = (
        ids_con_asistencia
        | ids_con_trabajos
        | ids_con_partidos
        | ids_con_ejercicios
    )

    turnos_realizados = (
        Entrenamiento.objects
        .filter(
            club=club,
            id__in=ids_turnos_con_actividad,
            no_se_entreno=False,
        )
        .count()
    )

    jugadores_datos = []

    for jugador in (
        Jugador.objects
        .filter(
            club=club,
            activo=True,
        )
        .order_by(
            "apellido",
            "nombre",
        )
    ):
        asistencias_jugador = asistencias_mes.filter(
            jugador=jugador,
        )

        total = asistencias_jugador.count()

        presentes = asistencias_jugador.filter(
            Q(estado="asistio")
            | Q(estado="tarde")
        ).count()

        tardes = asistencias_jugador.filter(
            estado="tarde",
        ).count()

        ausentes = asistencias_jugador.filter(
            estado="ausente",
        ).count()

        pendientes = asistencias_jugador.filter(
            estado="pendiente",
        ).count()

        porcentaje = (
            round(
                (presentes / total) * 100,
                1,
            )
            if total
            else 0
        )

        partidos_jugador = (
            PartidoTurno.objects
            .filter(
                entrenamiento__club=club,
                entrenamiento__fecha__range=[
                    inicio_mes,
                    fin_mes,
                ],
            )
            .filter(
                Q(jugador_1=jugador)
                | Q(jugador_2=jugador)
            )
            .select_related(
                "jugador_1",
                "jugador_2",
                "entrenamiento",
            )
            .prefetch_related("sets")
        )

        partidos = partidos_jugador.count()
        victorias = 0
        derrotas = 0

        for partido in partidos_jugador:
            sets_jugador_1 = partido.sets_jugador_1
            sets_jugador_2 = partido.sets_jugador_2

            if sets_jugador_1 == sets_jugador_2:
                continue

            if partido.jugador_1_id == jugador.id:
                if sets_jugador_1 > sets_jugador_2:
                    victorias += 1
                else:
                    derrotas += 1
            else:
                if sets_jugador_2 > sets_jugador_1:
                    victorias += 1
                else:
                    derrotas += 1

        trabajos = (
            TrabajoTurno.objects
            .filter(
                entrenamiento__club=club,
                entrenamiento__fecha__range=[
                    inicio_mes,
                    fin_mes,
                ],
            )
            .filter(
                Q(jugador_1=jugador)
                | Q(jugador_2=jugador)
            )
            .count()
        )

        observaciones = (
            ObservacionJugador.objects
            .filter(
                jugador=jugador,
                entrenamiento__club=club,
                entrenamiento__fecha__range=[
                    inicio_mes,
                    fin_mes,
                ],
            )
            .count()
        )

        jugadores_datos.append({
            "jugador": jugador,
            "total": total,
            "presentes": presentes,
            "tardes": tardes,
            "ausentes": ausentes,
            "pendientes": pendientes,
            "porcentaje": porcentaje,
            "partidos": partidos,
            "victorias": victorias,
            "derrotas": derrotas,
            "trabajos": trabajos,
            "observaciones": observaciones,
        })

    ranking_asistencia = sorted(
        jugadores_datos,
        key=lambda item: (
            item["porcentaje"],
            item["presentes"],
        ),
        reverse=True,
    )

    ranking_ausencias = sorted(
        jugadores_datos,
        key=lambda item: item["ausentes"],
        reverse=True,
    )

    ranking_partidos = sorted(
        jugadores_datos,
        key=lambda item: item["partidos"],
        reverse=True,
    )

    mejor_asistencia = (
        ranking_asistencia[0]
        if ranking_asistencia
        and ranking_asistencia[0]["total"] > 0
        else None
    )

    mas_ausencias = (
        ranking_ausencias[0]
        if ranking_ausencias
        and ranking_ausencias[0]["ausentes"] > 0
        else None
    )

    mas_partidos = (
        ranking_partidos[0]
        if ranking_partidos
        and ranking_partidos[0]["partidos"] > 0
        else None
    )

    total_presentes = sum(
        item["presentes"]
        for item in jugadores_datos
    )

    total_ausentes = sum(
        item["ausentes"]
        for item in jugadores_datos
    )

    total_tardes = sum(
        item["tardes"]
        for item in jugadores_datos
    )

    total_partidos = sum(
        item["partidos"]
        for item in jugadores_datos
    )

    resumen = {
        "jugadores_activos": len(jugadores_datos),
        "turnos_realizados": turnos_realizados,
        "turnos_no_entrenados": turnos_no_entrenados.count(),
        "dias_no_entrenados": (
            turnos_no_entrenados
            .values("fecha")
            .distinct()
            .count()
        ),
        "total_presentes": total_presentes,
        "total_ausentes": total_ausentes,
        "total_tardes": total_tardes,
        "total_partidos": total_partidos,
        "tiene_actividad_mes": (
            turnos_realizados > 0
            or turnos_no_entrenados.exists()
            or total_presentes > 0
            or total_ausentes > 0
            or total_tardes > 0
            or total_partidos > 0
        ),
        "motivo_principal": (
            motivos[0]["motivo"]
            if motivos
            else "-"
        ),
    }

    contexto = {
        "fecha_base": fecha_base,
        "inicio_mes": inicio_mes,
        "fin_mes": fin_mes,
        "mes_anterior": mes_anterior,
        "mes_siguiente": mes_siguiente,
        "resumen": resumen,
        "motivos": motivos,
        "turnos_no_entrenados": turnos_no_entrenados[:12],
        "ranking_asistencia": ranking_asistencia[:10],
        "ranking_ausencias": ranking_ausencias[:10],
        "ranking_partidos": ranking_partidos[:10],
        "mejor_asistencia": mejor_asistencia,
        "mas_ausencias": mas_ausencias,
        "mas_partidos": mas_partidos,
    }

    return render(
        request,
        "asistencia/dashboard_mensual.html",
        contexto,
    )



@login_required
def acerca(request):
    return render(request, "asistencia/acerca.html")

@login_required
def perfil(request):
    if request.method == "POST":
        form = PerfilForm(request.POST, instance=request.user)

        if form.is_valid():
            form.save()
            messages.success(request, "Perfil actualizado correctamente.")
            return redirect("perfil")
    else:
        form = PerfilForm(instance=request.user)

    return render(request, "asistencia/perfil.html", {
        "form": form,
    })
    
@login_required
@require_POST
def eliminar_observacion_jugador(request, observacion_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    observacion = get_object_or_404(
        ObservacionJugador.objects.select_related(
            "entrenamiento",
            "jugador",
        ),
        id=observacion_id,
        entrenamiento__club=club,
        jugador__club=club,
    )

    entrenamiento = observacion.entrenamiento
    jugador = observacion.jugador

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    observacion.delete()

    messages.success(
        request,
        f"Observación eliminada de {jugador}.",
    )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )
    
@login_required
def editar_observacion_jugador(request, observacion_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    observacion = get_object_or_404(
        ObservacionJugador.objects.select_related(
            "entrenamiento",
            "jugador",
        ),
        id=observacion_id,
        entrenamiento__club=club,
        jugador__club=club,
    )

    entrenamiento = observacion.entrenamiento

    if request.method == "POST" and turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    if request.method == "POST":
        form = ObservacionJugadorForm(
            request.POST,
            instance=observacion,
        )

        if form.is_valid():
            form.save()

            messages.success(
                request,
                "Observación actualizada correctamente.",
            )

            return redirect(
                "dia_turno",
                fecha_str=entrenamiento.fecha.isoformat(),
                turno_id=obtener_turno_id_para_url(entrenamiento),
            )
    else:
        form = ObservacionJugadorForm(
            instance=observacion,
        )

    return render(
        request,
        "asistencia/editar_observacion_jugador.html",
        {
            "form": form,
            "observacion": observacion,
            "entrenamiento": entrenamiento,
        },
    )
    
@login_required
def editar_partido_turno(request, partido_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    partido = get_object_or_404(
        PartidoTurno.objects.select_related(
            "entrenamiento",
            "jugador_1",
            "jugador_2",
        ),
        id=partido_id,
        entrenamiento__club=club,
        jugador_1__club=club,
        jugador_2__club=club,
    )

    entrenamiento = partido.entrenamiento

    if request.method == "POST" and turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    if request.method == "POST":
        partido_form = PartidoTurnoForm(
            request.POST,
            instance=partido,
            entrenamiento=entrenamiento,
        )

        sets_formset = SetPartidoFormSet(
            request.POST,
            instance=partido,
            prefix="sets",
        )

        if partido_form.is_valid() and sets_formset.is_valid():
            with transaction.atomic():
                partido = partido_form.save()
                sets = sets_formset.save(commit=False)

                numero_set = 1

                for set_partido in sets:
                    set_partido.partido = partido
                    set_partido.numero = numero_set
                    set_partido.save()
                    numero_set += 1

                for set_eliminado in sets_formset.deleted_objects:
                    set_eliminado.delete()

            messages.success(
                request,
                "El partido se actualizó correctamente.",
            )

            return redirect(
                "dia_turno",
                fecha_str=entrenamiento.fecha.isoformat(),
                turno_id=obtener_turno_id_para_url(entrenamiento),
            )

    else:
        partido_form = PartidoTurnoForm(
            instance=partido,
            entrenamiento=entrenamiento,
        )

        sets_formset = SetPartidoFormSet(
            instance=partido,
            prefix="sets",
        )

    return render(
        request,
        "asistencia/editar_partido_turno.html",
        {
            "partido": partido,
            "entrenamiento": entrenamiento,
            "partido_form": partido_form,
            "sets_formset": sets_formset,
        },
    )


@login_required
@require_POST
def eliminar_partido_turno(request, partido_id):
    club = obtener_club_usuario(request.user)

    if club is None:
        messages.error(
            request,
            "Tu usuario no está asociado a un club activo.",
        )
        return redirect("login")

    partido = get_object_or_404(
        PartidoTurno.objects.select_related(
            "entrenamiento",
            "jugador_1",
            "jugador_2",
        ),
        id=partido_id,
        entrenamiento__club=club,
        jugador_1__club=club,
        jugador_2__club=club,
    )

    entrenamiento = partido.entrenamiento

    if turno_bloqueado(entrenamiento):
        return redirigir_turno_bloqueado(
            request,
            entrenamiento,
        )

    partido.delete()

    messages.success(
        request,
        "El partido se eliminó correctamente.",
    )

    return redirect_dia_turno(
        entrenamiento,
        request.POST.get("volver_a", ""),
    )
    
