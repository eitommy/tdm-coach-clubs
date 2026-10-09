from io import BytesIO
from datetime import time

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from clubs.models import Club, PerfilUsuario, TurnoClub
from asistencia.models import (
    CategoriaEjercicio,
    CategoriaJugador,
    Ejercicio,
    Jugador,
    PagoJugador,
)


class MatrizPermisosRolesTests(TestCase):

    def setUp(self):
        self.club = Club.objects.create(
            nombre="Club Permisos",
            activo=True,
        )

        self.admin = User.objects.create_user(
            username="admin_permisos",
            password="TestPass123!",
            first_name="Admin",
            last_name="Permisos",
            is_active=True,
        )

        self.perfil_admin = PerfilUsuario.objects.create(
            usuario=self.admin,
            club=self.club,
            rol=PerfilUsuario.Rol.ADMIN,
            activo=True,
        )

        self.entrenador = User.objects.create_user(
            username="entrenador_permisos",
            password="TestPass123!",
            first_name="Entrenador",
            last_name="Permisos",
            is_active=True,
        )

        self.perfil_entrenador = PerfilUsuario.objects.create(
            usuario=self.entrenador,
            club=self.club,
            rol=PerfilUsuario.Rol.ENTRENADOR,
            activo=True,
        )

        self.turno = TurnoClub.objects.create(
            club=self.club,
            dia_semana=0,
            nombre="Turno prueba",
            hora_inicio=time(15, 0),
            hora_fin=time(17, 0),
            orden=1,
            activo=True,
        )

        self.categoria_jugador = CategoriaJugador.objects.create(
            club=self.club,
            nombre="Primera",
            orden=1,
            activo=True,
        )

        self.jugador = Jugador.objects.create(
            club=self.club,
            nombre="Jugador",
            apellido="Prueba",
            categoria=self.categoria_jugador,
            activo=True,
        )

        self.categoria_ejercicio = CategoriaEjercicio.objects.create(
            club=self.club,
            nombre="Categoría prueba",
            orden=1,
            activo=True,
        )

        self.ejercicio = Ejercicio.objects.create(
            club=self.club,
            nombre="Ejercicio prueba",
            categoria_config=self.categoria_ejercicio,
            activo=True,
        )

        PagoJugador.objects.create(
            jugador=self.jugador,
            anio=self.jugador.fecha_alta.year,
            mes=self.jugador.fecha_alta.month,
            pagado=True,
            fecha_pago=self.jugador.fecha_alta,
        )

    def _login_admin(self):
        self.client.force_login(self.admin)

    def _login_entrenador(self):
        self.client.force_login(self.entrenador)

    # ============================================================
    # ADMINISTRACIÓN GENERAL
    # ============================================================

    def test_entrenador_no_puede_entrar_a_configuracion_administrativa(self):
        self._login_entrenador()

        rutas = [
            ("configuracion_club", {}),
            ("lista_usuarios", {}),
            ("crear_usuario", {}),
            (
                "editar_usuario",
                {
                    "perfil_id": self.perfil_admin.id,
                },
            ),
        ]

        for nombre_url, kwargs in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url,
                        kwargs=kwargs or None,
                    )
                )

                self.assertEqual(
                    response.status_code,
                    403,
                    msg=(
                        f'El entrenador accedió a "{nombre_url}" '
                        f"con status {response.status_code}."
                    ),
                )

    def test_admin_puede_entrar_a_configuracion_administrativa(self):
        self._login_admin()

        rutas = [
            ("configuracion_club", {}),
            ("lista_usuarios", {}),
            ("crear_usuario", {}),
            (
                "editar_usuario",
                {
                    "perfil_id": self.perfil_entrenador.id,
                },
            ),
        ]

        for nombre_url, kwargs in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url,
                        kwargs=kwargs or None,
                    )
                )

                self.assertEqual(
                    response.status_code,
                    200,
                    msg=(
                        f'El admin no pudo acceder a "{nombre_url}". '
                        f"Status: {response.status_code}."
                    ),
                )

    # ============================================================
    # TURNOS - OPERATIVO PARA ENTRENADOR
    # ============================================================

    def test_entrenador_puede_gestionar_turnos(self):
        self._login_entrenador()

        rutas = [
            ("lista_turnos", {}),
            ("crear_turno", {}),
            (
                "editar_turno",
                {
                    "turno_id": self.turno.id,
                },
            ),
        ]

        for nombre_url, kwargs in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url,
                        kwargs=kwargs or None,
                    )
                )

                self.assertEqual(
                    response.status_code,
                    200,
                    msg=(
                        f'El entrenador no pudo acceder a "{nombre_url}". '
                        f"Status: {response.status_code}."
                    ),
                )

    def test_entrenador_puede_eliminar_turno_de_su_club(self):
        self._login_entrenador()

        turno = TurnoClub.objects.create(
            club=self.club,
            dia_semana=1,
            nombre="Turno eliminable",
            hora_inicio=time(18, 0),
            hora_fin=time(20, 0),
            orden=2,
            activo=True,
        )

        response = self.client.post(
            reverse(
                "eliminar_turno",
                kwargs={
                    "turno_id": turno.id,
                },
            )
        )

        self.assertEqual(
            response.status_code,
            302,
        )

        self.assertFalse(
            TurnoClub.objects.filter(
                pk=turno.id,
            ).exists()
        )

    # ============================================================
    # JUGADORES
    # ============================================================

    def test_entrenador_puede_gestionar_jugadores(self):
        self._login_entrenador()

        rutas = [
            ("lista_jugadores", {}),
            ("crear_jugador", {}),
            (
                "editar_jugador",
                {
                    "pk": self.jugador.id,
                },
            ),
            (
                "historial_jugador",
                {
                    "jugador_id": self.jugador.id,
                },
            ),
        ]

        for nombre_url, kwargs in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url,
                        kwargs=kwargs or None,
                    )
                )

                self.assertEqual(
                    response.status_code,
                    200,
                    msg=(
                        f'El entrenador no pudo acceder a "{nombre_url}". '
                        f"Status: {response.status_code}."
                    ),
                )

    # ============================================================
    # CATEGORÍAS DE JUGADOR - SOLO ADMIN
    # ============================================================

    def test_entrenador_no_puede_administrar_categorias_jugador(self):
        self._login_entrenador()

        rutas = [
            ("lista_categorias_jugador", {}),
            ("crear_categoria_jugador", {}),
            (
                "editar_categoria_jugador",
                {
                    "pk": self.categoria_jugador.id,
                },
            ),
        ]

        for nombre_url, kwargs in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url,
                        kwargs=kwargs or None,
                    )
                )

                self.assertEqual(
                    response.status_code,
                    403,
                    msg=(
                        f'El entrenador accedió a "{nombre_url}" '
                        f"con status {response.status_code}."
                    ),
                )

        response = self.client.post(
            reverse(
                "mover_categoria_jugador",
                kwargs={
                    "pk": self.categoria_jugador.id,
                    "direccion": "subir",
                },
            )
        )

        self.assertEqual(
            response.status_code,
            403,
        )

    def test_admin_puede_ver_categorias_jugador(self):
        self._login_admin()

        response = self.client.get(
            reverse(
                "lista_categorias_jugador"
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

    # ============================================================
    # CUOTAS - SOLO ADMIN
    # ============================================================

    def test_entrenador_no_puede_ver_resumen_cuotas(self):
        self._login_entrenador()

        response = self.client.get(
            reverse(
                "resumen_cuotas"
            )
        )

        self.assertEqual(
            response.status_code,
            403,
        )

    def test_entrenador_no_puede_exportar_cuotas(self):
        self._login_entrenador()

        response = self.client.get(
            reverse(
                "exportar_cuotas_excel"
            )
        )

        self.assertEqual(
            response.status_code,
            403,
        )

    def test_entrenador_no_puede_cambiar_estado_cuota(self):
        self._login_entrenador()

        pago = self.jugador.pagos.first()
        estado_antes = pago.pagado

        response = self.client.post(
            reverse(
                "cambiar_estado_cuota",
                kwargs={
                    "jugador_id": self.jugador.id,
                },
            ),
            data={
                "mes": pago.mes,
                "anio": pago.anio,
                "accion": "pendiente",
            },
        )

        self.assertEqual(
            response.status_code,
            403,
        )

        pago.refresh_from_db()

        self.assertEqual(
            pago.pagado,
            estado_antes,
        )

    def test_admin_puede_ver_resumen_y_exportar_cuotas(self):
        self._login_admin()

        response_resumen = self.client.get(
            reverse(
                "resumen_cuotas"
            )
        )

        self.assertEqual(
            response_resumen.status_code,
            200,
        )

        response_excel = self.client.get(
            reverse(
                "exportar_cuotas_excel"
            )
        )

        self.assertEqual(
            response_excel.status_code,
            200,
        )

        self.assertIn(
            "spreadsheetml.sheet",
            response_excel["Content-Type"],
        )

    # ============================================================
    # INFORMACIÓN ECONÓMICA EN HTML
    # ============================================================

    def test_lista_jugadores_entrenador_no_muestra_informacion_economica(self):
        self._login_entrenador()

        response = self.client.get(
            reverse(
                "lista_jugadores"
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        contenido = response.content.decode()

        self.assertNotIn(
            "Resumen cuotas",
            contenido,
        )

        self.assertNotIn(
            "Exportar Excel",
            contenido,
        )

        self.assertNotIn(
            "Marcar pagado",
            contenido,
        )

        self.assertNotIn(
            "Marcar pendiente",
            contenido,
        )

        self.assertNotIn(
            'id="filtro-cuota-jugador"',
            contenido,
        )

    def test_historial_entrenador_no_muestra_cuotas(self):
        self._login_entrenador()

        response = self.client.get(
            reverse(
                "historial_jugador",
                kwargs={
                    "jugador_id": self.jugador.id,
                },
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        contenido = response.content.decode()

        self.assertNotIn(
            "Historial de cuotas",
            contenido,
        )

        self.assertNotIn(
            'href="#cuotas-jugador"',
            contenido,
        )

    def test_historial_admin_si_muestra_cuotas(self):
        self._login_admin()

        response = self.client.get(
            reverse(
                "historial_jugador",
                kwargs={
                    "jugador_id": self.jugador.id,
                },
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        contenido = response.content.decode()

        self.assertIn(
            "Historial de cuotas",
            contenido,
        )

    # ============================================================
    # EJERCICIOS - OPERATIVO PARA ENTRENADOR
    # ============================================================

    def test_entrenador_puede_gestionar_catalogo_ejercicios(self):
        self._login_entrenador()

        rutas = [
            ("lista_ejercicios", {}),
            ("lista_categorias_ejercicio", {}),
            ("crear_categoria_ejercicio", {}),
            (
                "editar_categoria_ejercicio",
                {
                    "pk": self.categoria_ejercicio.id,
                },
            ),
            ("crear_ejercicio", {}),
            (
                "editar_ejercicio",
                {
                    "pk": self.ejercicio.id,
                },
            ),
        ]

        for nombre_url, kwargs in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url,
                        kwargs=kwargs or None,
                    )
                )

                self.assertEqual(
                    response.status_code,
                    200,
                    msg=(
                        f'El entrenador no pudo acceder a "{nombre_url}". '
                        f"Status: {response.status_code}."
                    ),
                )

    # ============================================================
    # REPORTES
    # ============================================================

    def test_entrenador_puede_ver_reportes(self):
        self._login_entrenador()

        rutas = [
            "reportes",
            "seguimiento_semanal",
            "dashboard_mensual",
            "resumen_dia",
        ]

        for nombre_url in rutas:
            with self.subTest(nombre_url=nombre_url):
                response = self.client.get(
                    reverse(
                        nombre_url
                    )
                )

                self.assertEqual(
                    response.status_code,
                    200,
                    msg=(
                        f'El entrenador no pudo acceder a "{nombre_url}". '
                        f"Status: {response.status_code}."
                    ),
                )

    # ============================================================
    # EXPORTACIÓN DE ASISTENCIA
    # ============================================================

    def test_entrenador_puede_exportar_asistencia(self):
        self._login_entrenador()

        response = self.client.get(
            reverse(
                "exportar_asistencia_excel"
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertIn(
            "spreadsheetml.sheet",
            response["Content-Type"],
        )

        workbook = load_workbook(
            BytesIO(
                response.content
            )
        )

        self.assertIn(
            "Resumen",
            workbook.sheetnames,
        )

        self.assertIn(
            "Asistencia detallada",
            workbook.sheetnames,
        )

    # ============================================================
    # EXPORTACIÓN DEL HISTORIAL
    # ============================================================

    def test_exportacion_historial_entrenador_no_incluye_cuotas(self):
        self._login_entrenador()

        response = self.client.get(
            reverse(
                "exportar_historial_jugador_excel",
                kwargs={
                    "jugador_id": self.jugador.id,
                },
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertIn(
            "spreadsheetml.sheet",
            response["Content-Type"],
        )

        workbook = load_workbook(
            BytesIO(
                response.content
            )
        )

        hojas_esperadas = {
            "Resumen",
            "Asistencias",
            "Partidos",
            "Trabajos",
            "Observaciones",
            "Ejercicios",
        }

        self.assertTrue(
            hojas_esperadas.issubset(
                set(
                    workbook.sheetnames
                )
            )
        )

        self.assertNotIn(
            "Cuotas",
            workbook.sheetnames,
        )

    def test_exportacion_historial_admin_si_incluye_cuotas(self):
        self._login_admin()

        response = self.client.get(
            reverse(
                "exportar_historial_jugador_excel",
                kwargs={
                    "jugador_id": self.jugador.id,
                },
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertIn(
            "spreadsheetml.sheet",
            response["Content-Type"],
        )

        workbook = load_workbook(
            BytesIO(
                response.content
            )
        )

        self.assertIn(
            "Cuotas",
            workbook.sheetnames,
        )